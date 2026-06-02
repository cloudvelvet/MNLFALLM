from pathlib import Path
import csv
import math
import random
from collections import defaultdict

from docx import Document


ROOT = Path(r"C:\chen_bauer_2024\MNLFALLM")
OUT = ROOT / "llm_dif_output"
PY_SEED = 20260531


def read_csv(path):
    with open(path, newline="", encoding="utf-8-sig") as f:
        return list(csv.DictReader(f))


def f3(x):
    if x is None or (isinstance(x, float) and math.isnan(x)):
        return ""
    return f"{float(x):.3f}".replace("0.", ".")


def average_precision(labels, scores):
    pairs = sorted(zip(scores, labels), key=lambda p: p[0], reverse=True)
    positives = sum(labels)
    if positives == 0:
        return float("nan")
    hits = 0
    total = 0.0
    for idx, (_, lab) in enumerate(pairs, start=1):
        if lab:
            hits += 1
            total += hits / idx
    return total / positives


def precision_at_k(labels, scores, k):
    pairs = sorted(zip(scores, labels), key=lambda p: p[0], reverse=True)
    top = pairs[: min(k, len(pairs))]
    if not top:
        return float("nan")
    return sum(lab for _, lab in top) / len(top)


def tie_random_metric(labels, scores, metric, reps=2000, seed=PY_SEED):
    rnd = random.Random(seed)
    vals = []
    by_score = defaultdict(list)
    for i, s in enumerate(scores):
        by_score[s].append(i)
    score_groups = sorted(by_score.keys(), reverse=True)
    for _ in range(reps):
        order = []
        for s in score_groups:
            group = by_score[s][:]
            rnd.shuffle(group)
            order.extend(group)
        ordered_labels = [labels[i] for i in order]
        ordered_scores = list(range(len(order), 0, -1))
        if metric == "ap":
            vals.append(average_precision(ordered_labels, ordered_scores))
        elif metric == "p5":
            vals.append(sum(ordered_labels[:5]) / min(5, len(ordered_labels)))
        elif metric == "p10":
            vals.append(sum(ordered_labels[:10]) / min(10, len(ordered_labels)))
    vals.sort()
    return sum(vals) / len(vals), vals[int(.025 * len(vals))], vals[int(.975 * len(vals)) - 1]


def group_metrics(rows, label_col="dif_label"):
    def as_label(value):
        text = str(value).strip().lower()
        if text in {"true", "t", "yes"}:
            return 1
        if text in {"false", "f", "no", ""}:
            return 0
        return int(float(value))

    labels = [as_label(r[label_col]) for r in rows]
    llm = [float(r["threshold_dif_probability_0_100"]) for r in rows]
    kw_col = "keyword_binary_score" if "keyword_binary_score" in rows[0] else "keyword_score_0_100"
    kw = [float(r[kw_col]) for r in rows]
    kw_ap, kw_ap_lo, kw_ap_hi = tie_random_metric(labels, kw, "ap")
    kw_p5, _, _ = tie_random_metric(labels, kw, "p5")
    kw_p10, _, _ = tie_random_metric(labels, kw, "p10")
    return {
        "n": len(rows),
        "pos": sum(labels),
        "prev": sum(labels) / len(rows) if rows else float("nan"),
        "llm_ap": average_precision(labels, llm),
        "llm_p5": precision_at_k(labels, llm, 5),
        "llm_p10": precision_at_k(labels, llm, 10),
        "kw_ap": kw_ap,
        "kw_ap_lo": kw_ap_lo,
        "kw_ap_hi": kw_ap_hi,
        "kw_p5": kw_p5,
        "kw_p10": kw_p10,
    }


def by_prompt(rows):
    out = {}
    for prompt in ["original", "strict_dif"]:
        sub = [r for r in rows if r["prompt_version"] == prompt]
        out[prompt] = group_metrics(sub)
    return out


def by_covariate(rows):
    labels = {
        "discrim_any": "차별 경험",
        "korean_c": "한국어 능력",
        "gender": "성별",
        "age_c": "연령",
        "income_c": "가구소득",
    }
    out = []
    for cov, name in labels.items():
        for prompt in ["original", "strict_dif"]:
            sub = [r for r in rows if r["covariate"] == cov and r["prompt_version"] == prompt]
            if sub:
                m = group_metrics(sub)
                out.append([name, "기본" if prompt == "original" else "엄격", m])
    return out


def set_paragraph(paragraph, text):
    paragraph.clear()
    paragraph.add_run(text)


def set_cell(cell, text):
    cell.text = str(text)


def rewrite_table(table, rows):
    while len(table.rows) < len(rows):
        table.add_row()
    while len(table.rows) > len(rows):
        tr = table.rows[-1]._tr
        tr.getparent().remove(tr)
    for r, row in enumerate(rows):
        while len(table.rows[r].cells) < len(row):
            # Word tables have fixed grid; this fallback should rarely be needed.
            table.add_column(1)
        for c, value in enumerate(row):
            set_cell(table.rows[r].cells[c], value)
        for c in range(len(row), len(table.rows[r].cells)):
            set_cell(table.rows[r].cells[c], "")


def main():
    doc_paths = list(Path("N:/").glob("개인/다음논문/LLM_DIF_full_manuscript_draft_psychometric_style.docx"))
    if not doc_paths:
        raise FileNotFoundError("Target manuscript not found on N:/개인/다음논문")
    doc_path = doc_paths[0]
    revised_path = doc_path.with_name(doc_path.stem + "_수정본.docx")

    main_rows = read_csv(OUT / "maps_llm_gemini_sensitivity_eval_joined_binary_keyword.csv")
    main_rows = [r for r in main_rows if r.get("parse_status", "ok") == "ok"]
    w6_rows = read_csv(OUT / "maps_llm_gemini_sensitivity_eval_joined_binary_keyword_w6.csv")
    w6_rows = [r for r in w6_rows if r.get("parse_status", "ok") == "ok"]
    overall = by_prompt(main_rows)
    overall_w6 = by_prompt(w6_rows)
    cov_rows = by_covariate(main_rows)

    # Threshold sensitivity from the same joined binary-keyword file.
    threshold_specs = [
        ("FDR<.05 only", lambda r: float(r["p_fdr"]) < .05),
        ("FDR<.05 + |beta|>=.10", lambda r: float(r["p_fdr"]) < .05 and abs(float(r["beta"])) >= .10),
        ("FDR<.05 + |beta|>=.20", lambda r: float(r["p_fdr"]) < .05 and abs(float(r["beta"])) >= .20),
        ("FDR<.05 + |beta|>=.30", lambda r: float(r["p_fdr"]) < .05 and abs(float(r["beta"])) >= .30),
    ]
    threshold_table_rows = [["\uc9c0\uc2dc\ubb38 \uc870\uac74", "\uc120\ubcc4 \uae30\uc900", "\uc591\uc131 \uc218", "\uc591\uc131 \ube44\uc728", "LLM AP", "\uc774\ubd84\ud615 \ud0a4\uc6cc\ub4dc AP", "LLM P@10", "\ud0a4\uc6cc\ub4dc P@10", "\uc8fc \ubd84\uc11d \uc591\uc131 \uc720\uc9c0\uc728"]]
    for prompt in ["original", "strict_dif"]:
        sub0 = [r for r in main_rows if r["prompt_version"] == prompt]
        base_positive = {
            (r["scale_id"], r["item_id"], r["covariate"])
            for r in sub0
            if float(r["p_fdr"]) < .05 and abs(float(r["beta"])) >= .20
        }
        for label, fn in threshold_specs:
            sub = [dict(r, tmp_label="1" if fn(r) else "0") for r in sub0]
            m = group_metrics(sub, label_col="tmp_label")
            current_positive = {
                (r["scale_id"], r["item_id"], r["covariate"])
                for r in sub0
                if fn(r)
            }
            retention = len(base_positive & current_positive) / len(base_positive) if base_positive else float("nan")
            threshold_table_rows.append([
                "기본" if prompt == "original" else "엄격",
                label,
                m["pos"],
                f3(m["prev"]),
                f3(m["llm_ap"]),
                f"{f3(m['kw_ap'])} [{f3(m['kw_ap_lo'])}, {f3(m['kw_ap_hi'])}]",
                f3(m["llm_p10"]),
                f3(m["kw_p10"]),
                f3(retention),
            ])

    # Cluster-robust sensitivity using binary keyword, if the joined file exists.
    cr_path = OUT / "maps_llm_gemini_sensitivity_eval_joined_cluster_robust_full.csv"
    cr_rows = []
    if cr_path.exists():
        cr_source = read_csv(cr_path)
        kw_map = {
            (r["scale_id"], r["item_id"], r["covariate"]): r.get("keyword_binary_score", r.get("keyword_score_0_100"))
            for r in main_rows
        }
        for r in cr_source:
            key = (r["scale_id"], r["item_id"], r["covariate"])
            if key in kw_map and r.get("parse_status", "ok") == "ok":
                r["keyword_binary_score"] = kw_map[key]
                cr_rows.append(r)
    cluster_rows = [["지시문 조건", "분석 조합 수", "선별 양성 수", "선별 양성 비율", "LLM AP", "이분형 키워드 AP", "LLM P@10", "키워드 P@10"]]
    if cr_rows:
        for prompt in ["original", "strict_dif"]:
            sub = [r for r in cr_rows if r["prompt_version"] == prompt]
            m = group_metrics(sub, label_col="dif_label_cluster")
            cluster_rows.append(["기본" if prompt == "original" else "엄격", m["n"], m["pos"], f3(m["prev"]), f3(m["llm_ap"]), f3(m["kw_ap"]), f3(m["llm_p10"]), f3(m["kw_p10"])])

    doc = Document(doc_path)

    # Abstracts.
    set_paragraph(doc.paragraphs[4],
        "본 연구는 다문화청소년패널(MAPS) 2기 자료의 청소년 및 보호자 문항을 대상으로, 생성형 대규모 언어모형이 DIF 검토를 위한 문항-공변량 후보를 사전에 우선순위화할 수 있는지 탐색하였다. 분석 단위는 문항-공변량 조합이며, Gemini 2.5 Flash에는 문항 문장, 구인 맥락, 응답 범주, 공변량 정의를 제공하되 경험적 DIF 선별 결과, p값, 효과크기, 문항 통계량은 제공하지 않았다. LLM 산출물은 DIF 판정이 아니라 후속 검토가 필요한 후보 가설로 정의하였다.")
    set_paragraph(doc.paragraphs[5],
        "LLM 점수는 이분형 키워드 비교 기준 및 잠정적 ordinal DIF 선별 결과와 비교하였다. 이분형 키워드 기준은 문항 텍스트에 공변량 관련 핵심어가 하나 이상 포함되면 100점, 그렇지 않으면 0점을 부여하는 방식으로 구성하였다. 동점이 많은 기준의 특성을 고려하여 키워드 AP와 P@k는 무작위 동점 해소를 반복한 평균으로 보고하였다. 주 분석에서 기본 지시문 LLM의 AP는 .382, 엄격 지시문 LLM의 AP는 .405였고, 이분형 키워드 기준의 AP 평균은 각각 .375와 .384였다. 전체 AP에서는 LLM이 키워드 기준을 소폭 상회했으나, 이 우위는 공변량과 분석 조건 전반에서 일관되지 않았다. 6차년도 단일 wave 민감도 분석에서는 키워드 기준의 전체 AP가 LLM보다 높았다.")
    set_paragraph(doc.paragraphs[6],
        "결과는 LLM을 DIF 판정 도구로 사용하기에는 근거가 부족하지만, 문항-공변량 조합에 대한 후보 설명을 생성하고 검토 순서를 제안하는 보조 도구로는 제한적 가능성이 있음을 보여준다. 특히 LLM의 가치는 단순 예측 성능보다 판단 근거가 어떤 방식으로 문항 표현, 응답 과정, 문화·언어 맥락, 또는 부적절한 일반화와 연결되는지를 드러내는 데 있다. 따라서 LLM 기반 DIF 후보 생성은 키워드 기준, 경험적 선별, 전문가 검토, 후속 심리측정 모형과 결합될 때에만 해석 가능하다.")
    set_paragraph(doc.paragraphs[9],
        "This study examined whether a generative large language model can prioritize item-covariate combinations for subsequent differential item functioning (DIF) review using youth and caregiver items from the second cohort of the Multicultural Adolescents Panel Study. The unit of analysis was the item-covariate combination. Gemini 2.5 Flash received item text, construct context, response categories, and covariate definitions, but not empirical DIF screening results, p values, effect sizes, or item statistics. LLM outputs were treated as candidate hypotheses rather than DIF decisions.")
    set_paragraph(doc.paragraphs[10],
        "LLM scores were compared with a binary keyword benchmark and provisional ordinal DIF screening results. The binary keyword benchmark assigned 100 points when an item contained at least one prespecified covariate-related keyword and 0 otherwise. Because the benchmark produced many ties, keyword AP and P@k were reported as averages over repeated random tie-breaking. In the pooled main analysis, AP was .382 for the original prompt and .405 for the strict DIF prompt, whereas the binary keyword benchmark yielded mean AP values of .375 and .384, respectively. Thus, the LLM slightly exceeded the binary keyword benchmark in overall AP, but this advantage was not consistent across covariates or sensitivity analyses. In the single-wave Wave 6 analysis, the keyword benchmark had higher overall AP than the LLM.")
    set_paragraph(doc.paragraphs[11],
        "These findings do not support the use of LLMs as DIF decision tools. They suggest a more limited role: generating reviewable candidate explanations and helping organize the front end of DIF source investigation. LLM-based workflows remain interpretable only when combined with lexical benchmarks, empirical screening, expert review, and subsequent psychometric validation.")

    # Methods: keyword benchmark.
    set_paragraph(doc.paragraphs[46],
        "LLM 산출물이 문항의 표면 어휘를 넘어 추가적인 정보를 제공하는지 확인하기 위해 이분형 키워드 비교 기준을 구성하였다. 이 기준은 문항 텍스트에 공변량 관련 핵심어가 하나 이상 포함되는지를 기준으로 점수화하였다. LLM 출력, 경험적 선별 결과, 응답 분포, 문항 통계량은 사용하지 않았다. 다만 이 기준은 무작위 기준이 아니라 연구자가 사전에 구성한 어휘 기반 기준이므로, 차별 경험이나 한국어 능력처럼 표면 단서가 강한 공변량에서는 상당히 강한 비교 기준으로 작동할 수 있다.")
    set_paragraph(doc.paragraphs[47],
        "차별 경험에는 다른 대우, 편견, 무시, 위축, 사회적 지위, 따돌, 못살게, 외국, 욕, 놀림, 소문을 포함하였다. 한국어 능력에는 한국어, 한국문화, 한국 사람, 한국사람, 한국에, 한국의, 모국, 외국, 문화, 언어를 포함하였다. 가구소득, 성별, 연령에 대해서도 각각 경제적 조건, 또래 관계와 외모, 발달 단계와 진로에 관련된 핵심어를 지정하였다. 키워드 목록은 LLM 결과와 경험적 선별 결과를 보지 않은 상태에서 구성하였다.")
    set_paragraph(doc.paragraphs[48],
        "점수화는 의도적으로 단순하게 하였다. 각 문항-공변량 조합에 대해 해당 공변량의 키워드가 하나 이상 발견되면 100점, 발견되지 않으면 0점을 부여하였다. 문자열 탐색은 고정 문자열 일치 방식으로 수행하였고, 형태소 분석, 동의어 확장, TF-IDF, 문항 빈도 가중은 사용하지 않았다. 이분형 기준은 동점이 많이 발생하므로, 결과표에서는 키워드 기준의 AP와 P@k를 무작위 동점 해소를 2,000회 반복한 평균으로 제시하였다.")

    # Results.
    set_paragraph(doc.paragraphs[69],
        "표 4는 주 분석에서 LLM 점수와 이분형 키워드 비교 기준의 우선순위화 성능을 비교한 결과이다. 기본 지시문 조건에서 분석 가능한 조합은 486개였고, 이 중 잠정적 경험 DIF 후보는 128개였다. 엄격 지시문 조건에서는 482개 조합 중 127개가 잠정적 경험 DIF 후보로 분류되었다. 양성 비율은 두 조건 모두 약 .263이었다.")
    set_paragraph(doc.paragraphs[70],
        "기본 지시문 조건에서 LLM의 AP는 .382였고, 이분형 키워드 기준의 평균 AP는 .375였다. 엄격 지시문 조건에서는 LLM의 AP가 .405, 키워드 기준의 평균 AP가 .384였다. 즉 전체 AP만 보면 LLM은 이분형 키워드 기준을 소폭 상회하였다. 그러나 차이는 크지 않았고, 상위 5개 및 상위 10개 후보의 정밀도에서는 지시문 조건에 따라 양상이 달랐다.")
    set_paragraph(doc.paragraphs[71],
        "따라서 주 분석 결과는 LLM이 단순 어휘 기준보다 압도적으로 우수하다는 결론을 지지하지 않는다. 오히려 LLM의 추가 가치는 제한적이며, 키워드 기준처럼 투명한 비교 기준을 함께 두어야 그 의미를 해석할 수 있음을 보여준다.")
    set_paragraph(doc.paragraphs[72],
        "기본 지시문에서 LLM의 P@5와 P@10은 모두 .600이었다. 엄격 지시문에서는 AP가 높아졌지만 P@5는 .400, P@10은 .300으로 낮아졌다. 이는 엄격한 DIF 구분 지시문이 전체 순위에서는 일부 개선을 보였더라도, 최상위 후보 목록의 질을 안정적으로 개선하지는 못했음을 의미한다.")
    set_paragraph(doc.paragraphs[73],
        "이 결과는 LLM 기반 후보 생성에서 평균적인 순위 지표와 실제 검토 목록의 유용성이 항상 같은 방향으로 움직이지 않을 수 있음을 보여준다.")

    set_paragraph(doc.paragraphs[75],
        "공변량별 결과는 LLM의 성능이 공변량의 성격에 따라 달라짐을 보여준다(표 5). 차별 경험에서는 기본 지시문에서 LLM AP(.612)가 키워드 기준(.620)과 거의 같았고, 엄격 지시문에서는 LLM AP(.681)가 키워드 기준(.621)을 넘어섰다. 이는 차별 경험처럼 문항 내용과 공변량의 의미가 직접 맞물리는 경우, 엄격한 DIF 구분 지시문이 일부 도움이 될 수 있음을 시사한다.")
    set_paragraph(doc.paragraphs[76],
        "반면 한국어 능력에서는 두 지시문 조건 모두 키워드 기준이 LLM보다 높았다. 한국어, 문화, 외국과 같은 표면 어휘가 문항에 직접 나타나는 경우가 많았기 때문에, 단순 어휘 기준이 강한 비교 기준으로 작동한 것으로 볼 수 있다. 연령에서도 키워드 기준이 LLM보다 높았다.")
    set_paragraph(doc.paragraphs[77],
        "성별에서는 기본 지시문 LLM이 키워드 기준보다 약간 높았으나, 엄격 지시문에서는 그 차이가 사라졌다. 가구소득은 잠정적 경험 후보가 세 개뿐이어서 AP가 매우 불안정하므로 강한 해석을 피해야 한다.")
    set_paragraph(doc.paragraphs[78],
        "전체적으로 공변량별 분석은 LLM의 유용성이 단일 평균값으로 요약되기 어렵다는 점을 보여준다. LLM은 일부 공변량에서 키워드 기준을 보완할 수 있지만, 표면 단서가 강한 공변량에서는 단순 키워드 기준을 안정적으로 넘어서지 못했다.")

    # Sensitivity sections.
    set_paragraph(doc.paragraphs[100],
        "잠정적 경험 DIF 선별 기준을 달리했을 때도 결론의 방향은 크게 달라지지 않았다(표 10). 기준을 완화하거나 강화하면 양성 조합 수와 양성 비율은 변하지만, LLM과 키워드 기준의 상대적 성능은 조건에 따라 달라졌다. 특히 효과크기 기준을 강하게 적용할수록 양성 조합 수가 줄어들기 때문에 AP의 안정성도 함께 낮아진다. 따라서 본 연구의 결과는 특정 임계값 하나에 의존한 확정적 결론이라기보다, 여러 잠정적 기준 아래에서 LLM의 제한적이고 조건부적인 우선순위화 가능성을 확인한 결과로 해석해야 한다.")
    set_paragraph(doc.paragraphs[103],
        "반복측정 자료의 응답자 내 의존성을 고려하기 위해 cluster-robust 표준오차를 사용한 선별 결과도 민감도 분석으로 검토하였다(표 11). 이 분석에서도 LLM과 이분형 키워드 기준의 차이는 크지 않았으며, 결과의 방향은 지시문 조건과 선별 기준에 따라 달라졌다. 따라서 cluster-robust 분석은 LLM의 일관된 우위를 뒷받침하기보다는, 경험적 label 자체가 모형 선택과 표준오차 처리 방식에 민감할 수 있음을 보여주는 보조 결과로 보는 것이 적절하다.")
    set_paragraph(doc.paragraphs[108],
        "6차년도 단일 wave만 사용한 민감도 분석에서는 주 분석과 다른 양상이 나타났다(표 12). 기본 지시문 조건에서 LLM AP는 .208, 이분형 키워드 기준의 평균 AP는 .248이었다. 엄격 지시문 조건에서도 LLM AP는 .201, 키워드 기준의 평균 AP는 .254였다. 즉 6차년도 단일 wave에서는 전체 AP 기준으로 키워드 기준이 LLM보다 높았다.")
    set_paragraph(doc.paragraphs[111],
        "이 결과는 1-5차년도 pooled 분석에서 나타난 LLM의 소폭 우위가 단일 wave 분석으로 일반화되지 않음을 보여준다. 따라서 본 연구의 결론은 LLM의 안정적 예측 우위가 아니라, 분석 조건에 따라 달라지는 제한적 후보 생성 가능성으로 제시되어야 한다.")

    set_paragraph(doc.paragraphs[113],
        "본 연구는 MAPS 문항-공변량 조합을 대상으로, 경험적 DIF 선별 결과를 보지 않은 LLM이 후속 검토를 위한 후보 가설을 생성하고 우선순위화할 수 있는지 평가하였다. 결과는 네 가지로 요약된다.")
    set_paragraph(doc.paragraphs[114],
        "첫째, LLM은 대부분의 문항-공변량 조합에 대해 구조화된 후보 가설과 판단 근거를 생성할 수 있었다. 이는 LLM이 문항 문장, 구인 맥락, 응답 범주, 공변량 정의를 함께 읽고 검토 가능한 설명을 산출할 수 있음을 보여준다. 그러나 이 결과는 LLM이 DIF를 판정할 수 있음을 의미하지 않는다. LLM의 산출물은 경험적 검증과 전문가 검토를 기다리는 후보 가설에 머물러야 한다.")
    set_paragraph(doc.paragraphs[115],
        "둘째, LLM의 우선순위화 성능은 제한적이었다. 1-5차년도 pooled 주 분석의 전체 AP에서는 LLM이 이분형 키워드 기준을 소폭 상회하였다. 그러나 이 우위는 크지 않았고, 공변량별 분석과 6차년도 단일 wave 민감도 분석에서는 일관되게 유지되지 않았다. 따라서 본 연구의 결과는 LLM의 일반적 우수성을 보여주기보다, LLM 산출물이 투명한 어휘 기준과 나란히 평가되어야 함을 보여준다.")
    set_paragraph(doc.paragraphs[116],
        "셋째, 성능은 공변량의 성격에 따라 달랐다. 차별 경험에서는 엄격한 DIF 구분 지시문이 LLM 성능을 높였지만, 한국어 능력과 연령에서는 키워드 기준이 더 높았다. 성별 결과는 일부 조건에서 LLM이 높았으나 차이가 작고, 가구소득은 양성 후보 수가 적어 해석이 제한된다. 이는 LLM이 모든 공변량에 대해 같은 방식으로 작동하지 않으며, 공변량의 의미 구조와 문항의 표면 단서가 결과를 크게 좌우한다는 점을 시사한다.")
    set_paragraph(doc.paragraphs[117],
        "넷째, 엄격한 DIF 구분 지시문은 부분적으로만 도움이 되었다. 전체 AP는 기본 지시문보다 높아졌지만, 상위 후보의 정밀도는 낮아졌고 6차년도 민감도 분석에서도 우위가 유지되지 않았다. 지시문은 LLM의 설명 방식을 조정할 수 있지만, 그 자체로 안정적인 후보 목록을 보장하지는 못한다.")
    set_paragraph(doc.paragraphs[118],
        "이러한 결과는 LLM을 DIF 연구에 사용할 때 기대와 역할을 분명히 제한해야 함을 보여준다. LLM은 경험적 DIF의 대체물이 아니며, 문항 편향이나 불공정성을 판정하는 도구도 아니다. LLM의 잠재적 가치는 후보 공간을 정리하고, 문항과 공변량 사이의 가능한 의미 연결을 명시화하며, 그 설명이 어디에서 심리측정 검증과 맞물리고 어디에서 어긋나는지를 드러내는 데 있다.")
    set_paragraph(doc.paragraphs[124],
        "셋째, 키워드 비교 기준은 투명하지만 중립적인 기준은 아니다. 본 연구에서는 임의 가중치를 제거하고, 키워드 포함 여부만을 사용하는 이분형 기준을 사용하였다. 그럼에도 키워드 목록은 연구자의 판단에 의해 구성되었으며, 차별 경험이나 한국어 능력처럼 표면 어휘가 강한 공변량에서는 매우 강한 비교 기준으로 작동할 수 있다. 또한 이분형 점수는 동점이 많기 때문에 동점 처리 방식에 따라 AP와 P@k가 달라질 수 있다. 본 연구는 무작위 동점 해소 반복 평균을 보고했지만, 후속 연구에서는 키워드 사전의 사전등록, 독립 연구자 검토, 형태소 기반 대안 기준을 함께 검토할 필요가 있다.")
    set_paragraph(doc.paragraphs[126],
        "결론적으로, 본 연구는 LLM을 DIF 판정 도구로 지지하지 않는다. 주 분석에서 LLM은 이분형 키워드 기준을 전체 AP에서 소폭 상회했지만, 그 차이는 작았고 공변량별·시점별 분석에서 안정적으로 반복되지 않았다. 따라서 본 연구의 핵심 결론은 LLM이 키워드 기준을 대체할 수 있다는 것이 아니라, LLM 산출물을 키워드 기준과 경험적 선별 결과에 비추어 제한적으로 평가해야 한다는 데 있다.")
    set_paragraph(doc.paragraphs[127],
        "LLM은 답을 내리는 도구가 아니라 검토할 후보를 제안하는 도구다. 그 후보가 실제 문항 기능 차이와 관련되는지는 별도의 경험적 분석과 전문가 검토, 후속 심리측정 모형을 통해 확인되어야 한다. 이 역할 분담을 분명히 할 때, LLM은 DIF 원천 탐색의 앞단에서 후보 설명을 생성하고 그 한계를 함께 드러내는 보조 절차로 활용될 수 있다.")

    # Tables.
    rewrite_table(doc.tables[3], [
        ["공변량", "사용 키워드", "점수화 규칙", "구성 시점/blinding"],
        ["차별 경험", "다른 대우, 편견, 무시, 위축, 사회적 지위, 따돌, 못살게, 외국, 욕, 놀림, 소문", "키워드 포함=100점, 미포함=0점", "screening 결과와 LLM 결과 미사용"],
        ["한국어 능력", "한국어, 한국문화, 한국 사람, 한국사람, 한국에, 한국의, 모국, 외국, 문화, 언어", "키워드 포함=100점, 미포함=0점", "screening 결과와 LLM 결과 미사용"],
        ["가구소득", "경제, 형편, 물건, 장소, 제공, 대학, 회사, 지위, 건강, 병원", "키워드 포함=100점, 미포함=0점", "screening 결과와 LLM 결과 미사용"],
        ["성별", "외모, 이성친구, 신체적 특징, 친구, 따돌림, 소문, 욕설, 놀림", "키워드 포함=100점, 미포함=0점", "screening 결과와 LLM 결과 미사용"],
        ["연령", "진학, 진로, 미래, 부모역할, 자녀, 아이, 대학, 회사, 체류, 비자", "키워드 포함=100점, 미포함=0점", "screening 결과와 LLM 결과 미사용"],
    ])

    table4 = [["지시문 조건", "분석 조합 수", "선별 양성 수", "양성 비율", "LLM AP", "키워드 AP", "LLM P@5", "키워드 P@5", "LLM P@10", "키워드 P@10"]]
    for prompt in ["original", "strict_dif"]:
        label = "기본 지시문" if prompt == "original" else "엄격 지시문"
        m = overall[prompt]
        table4.append([label, m["n"], m["pos"], f3(m["prev"]), f3(m["llm_ap"]), f3(m["kw_ap"]), f3(m["llm_p5"]), f3(m["kw_p5"]), f3(m["llm_p10"]), f3(m["kw_p10"])])
    rewrite_table(doc.tables[4], table4)

    table5 = [["공변량", "지시문 조건", "분석 조합 수", "선별 양성 수", "양성 비율", "LLM AP", "키워드 AP"]]
    for name, prompt_label, m in cov_rows:
        table5.append([name, prompt_label, m["n"], m["pos"], f3(m["prev"]), f3(m["llm_ap"]), f3(m["kw_ap"])])
    rewrite_table(doc.tables[5], table5)

    rewrite_table(doc.tables[10], threshold_table_rows)
    if len(cluster_rows) > 1:
        rewrite_table(doc.tables[11], cluster_rows)

    table12 = [["지시문 조건", "분석 조합 수", "선별 양성 수", "양성 비율", "LLM AP", "키워드 AP", "LLM P@5", "키워드 P@5", "LLM P@10", "키워드 P@10"]]
    for prompt in ["original", "strict_dif"]:
        label = "기본 지시문" if prompt == "original" else "엄격 지시문"
        m = overall_w6[prompt]
        table12.append([label, m["n"], m["pos"], f3(m["prev"]), f3(m["llm_ap"]), f3(m["kw_ap"]), f3(m["llm_p5"]), f3(m["kw_p5"]), f3(m["llm_p10"]), f3(m["kw_p10"])])
    rewrite_table(doc.tables[12], table12)

    doc.save(revised_path)
    print(revised_path)


if __name__ == "__main__":
    main()
