import os
import re
from copy import deepcopy

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor
from docx.text.paragraph import Paragraph


SRC = r"N:\개인\다음논문\LLM_DIF_full_manuscript_draft_psychometric_style_수정본_방법결과반영_글자크기정리_SSCI보강_리뷰반영.docx"
OUT = r"N:\개인\다음논문\LLM_DIF_full_manuscript_draft_psychometric_style_정리본.docx"


TITLE = "생성형 언어모형은 순서형 DIF 후보를 키워드 기준보다 잘 우선순위화하는가?"
SUBTITLE = "다문화청소년패널 문항-공변량 조합의 벤치마크와 경계조건 분석"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=90, bottom=90, end=90):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = tr_pr.find(qn("w:tblHeader"))
    if tbl_header is None:
        tbl_header = OxmlElement("w:tblHeader")
        tr_pr.append(tbl_header)
    tbl_header.set(qn("w:val"), "true")


def set_paragraph_text(p, text):
    p.clear()
    r = p.add_run(text)
    r.font.name = "맑은 고딕"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
    return r


def insert_paragraph_before_element(element, parent, text):
    new_p = OxmlElement("w:p")
    element.addprevious(new_p)
    p = Paragraph(new_p, parent)
    r = p.add_run(text)
    r.font.name = "맑은 고딕"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
    return p


def iter_block_items(doc):
    for child in doc.element.body.iterchildren():
        tag = child.tag.split("}")[-1]
        if tag == "p":
            yield "p", Paragraph(child, doc)
        elif tag == "tbl":
            yield "tbl", child


def normalize_table_numbers(doc):
    # Add the missing caption for the main performance table before the first table in 3.2.
    blocks = list(iter_block_items(doc))
    for i, (kind, obj) in enumerate(blocks):
        if kind == "p" and obj.text.strip().startswith("기본 지시문 조건에서 LLM의 AP는 .382"):
            # next table is the whole-performance table
            for j in range(i + 1, len(blocks)):
                if blocks[j][0] == "tbl":
                    prev_text = blocks[j - 1][1].text.strip() if blocks[j - 1][0] == "p" else ""
                    if not prev_text.startswith("표 5."):
                        cap = insert_paragraph_before_element(blocks[j][1], doc._body, "표 5. 전체 우선순위화 성능")
                        cap.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    break
            break

    # Rename captions after the newly inserted Table 5.
    replacements = {
        "표 6. 공변량별 AP": "표 6. 공변량별 AP",
        "표 6-1. 응답자 유형별 LLM과 키워드 기준의 AP 비교": "표 7. 응답자 유형별 LLM과 키워드 기준의 AP 비교",
        "표 6-2. LLM 점수 구간별 잠정적 선별 양성률": "표 8. LLM 점수 구간별 잠정적 선별 양성률",
        "표 7. 지시문 조건 간 민감도": "표 9. 지시문 조건 간 민감도",
        "표 8. LLM 판단 근거의 자동 코딩 기반 설명 양식 요약": "표 10. LLM 판단 근거의 자동 코딩 기반 설명 양식 요약",
        "표 9. 계획된 수동 검증 표본 및 코딩 항목": "표 11. 계획된 수동 검증 표본 및 코딩 항목",
        "표 10. 대표 후보 사례": "표 12. 대표 후보 사례",
        "표 11. 선별 양성 기준 민감도 분석": "표 13. 선별 양성 기준 민감도 분석",
        "표 12. 응답자 단위 cluster-robust 민감도 분석": "표 14. 응답자 단위 cluster-robust 민감도 분석",
        "표 13. 6차년도 단일 wave 분석": "표 15. 6차년도 단일 wave 분석",
        "표 14. LLM과 키워드 기준 AP 차이에 대한 bootstrap 및 permutation 결과": "표 16. LLM과 키워드 기준 AP 차이에 대한 bootstrap 및 permutation 결과",
        "표 15. TF-IDF 의미 유사도 기준과 LLM의 우선순위화 성능": "표 17. TF-IDF 의미 유사도 기준과 LLM의 우선순위화 성능",
        "표 16. LLM 점수와 키워드 기준의 사분면별 경험적 선별 양상": "표 18. LLM 점수와 키워드 기준의 사분면별 경험적 선별 양상",
        "표 17. 사분면별 focal 사례의 random-intercept ordinal probe 요약": "표 19. 사분면별 focal 사례의 random-intercept ordinal probe 요약",
    }
    for p in doc.paragraphs:
        txt = p.text.strip()
        if txt in replacements:
            set_paragraph_text(p, replacements[txt])

    # Text references after Table 4.
    ref_replacements = [
        ("표 4는 주 분석", "표 5는 주 분석"),
        ("(표 5)", "(표 6)"),
        ("표 6-1", "표 7"),
        ("표 6-2", "표 8"),
        ("표 7)", "표 9)"),
        ("표 8)", "표 10)"),
        ("표 9)", "표 11)"),
        ("표 10)", "표 12)"),
        ("표 11)", "표 13)"),
        ("표 12)", "표 14)"),
        ("표 13)", "표 15)"),
        ("표 14)", "표 16)"),
        ("표 15)", "표 17)"),
        ("표 16)", "표 18)"),
        ("표 17)", "표 19)"),
    ]
    for p in doc.paragraphs:
        txt = p.text
        new = txt
        for a, b in ref_replacements:
            new = new.replace(a, b)
        if new != txt:
            set_paragraph_text(p, new)


def apply_text_polish(doc):
    # B3/B4 language and "representative" cleanup.
    for p in doc.paragraphs:
        txt = p.text.strip()
        if txt.startswith("언어모형을 심리측정 문항 생성과 검토에 활용하려는 시도"):
            set_paragraph_text(
                p,
                "언어모형을 심리측정 문항 생성과 검토에 활용하려는 시도는 이미 진행되고 있다. Transformer 기반 언어모형은 자동 문항 생성과 문항 검토에 활용될 가능성이 논의되어 왔으며(Attali et al., 2022; Hommel et al., 2022), AI가 생성한 문항 역시 인간이 작성한 문항과 마찬가지로 공정성과 DIF 검토를 필요로 한다는 점도 제기되었다(Belzak et al., 2023). DIF와 더 직접적으로 관련하여 Maeda와 Lu(2025)는 encoder 기반 transformer 모형을 사용해 문항 텍스트에서 경험적 DIF를 예측하고, 설명 가능한 인공지능 기법으로 관련 단어를 탐색하였다. Maeda와 Lu(2025)의 접근은 경험적 DIF 레이블이 있는 문항 텍스트로 모형을 훈련해 예측하는 방식인 반면, 본 연구는 경험적 선별 결과를 제공하지 않은 상태에서 생성형 LLM이 문항-공변량 의미 관계를 언어적으로 추론하고 후보를 우선순위화할 수 있는지를 평가한다.",
            )
        elif txt.startswith("주요 평가지표는 average precision(AP)이다."):
            set_paragraph_text(
                p,
                "주요 평가지표는 average precision(AP)이다. AP는 선별 양성 조합이 점수 순위의 상위에 얼마나 집중되는지를 평가하며, 양성 비율이 낮은 후보 우선순위화 문제에서 유용한 precision-recall 기반 지표이다(Davis & Goadrich, 2006). 본 연구의 목적은 최종 이분 분류가 아니라 전문가 검토 전 후보 우선순위화이므로, 정확도보다 AP를 주 지표로 사용하였다. AP는 양성 비율의 영향을 받기 때문에 결과표에는 선별 양성 비율을 무작위 순위화의 기준값으로 함께 제시하였다. 보조 지표로 precision@5와 precision@10을 보고하였다. LLM과 키워드 기준의 차이에 대한 불확실성 평가는 2.7절의 cluster bootstrap과 paired permutation test에서 별도로 다룬다.",
            )
        elif txt.startswith("마지막으로 LLM 점수와 키워드 기준이 수렴하거나 불일치하는 조합"):
            set_paragraph_text(
                p,
                "마지막으로 LLM 점수와 키워드 기준이 수렴하거나 불일치하는 조합의 경험적 양상을 검토하기 위해 사분면 분석을 수행하였다. LLM-high와 LLM-low의 기준은 자료에서 성능을 최적화하여 정한 값이 아니다. 본 연구에서는 LLM 출력이 0-100 척도로 산출된다는 점을 고려하여, 70점 이상을 높은 후보군, 30점 이하를 낮은 후보군으로 구분하는 해석적 절단점으로 사용하였다. 즉, 이 기준은 최적 분류 threshold가 아니라 LLM과 키워드 기준의 수렴 및 불일치 양상을 기술하기 위한 사전적 분석 기준이다. Keyword-high는 해당 공변량 키워드가 하나 이상 포함된 경우로 정의하였다. 이에 따라 LLM-high/keyword-low, keyword-high/LLM-low, both-high, both-low의 네 영역을 구성하였다.",
            )
        elif txt.startswith("각 영역에 대해서는 주 선별 기준과 cluster-robust 선별 기준"):
            set_paragraph_text(
                p,
                "각 영역에 대해서는 주 선별 기준과 cluster-robust 선별 기준의 양성 비율을 비교하였다. 추가로, 각 지시문 조건에서 네 사분면별로 1개 문항-공변량 조합을 선정하여 총 8개 조합에 대해 random-intercept ordinal probe를 수행하였다. 이 사례들은 각 사분면 전체를 대표한다고 가정하기 위한 것이 아니라, 사분면별 후보가 응답자별 무선절편을 포함한 모형에서도 같은 방향의 신호를 보이는지 점검하기 위한 focal 사례로 선정하였다. 선정은 경험적 선별 결과를 새로 최적화하는 방식이 아니라, 해당 사분면에 속하면서 분석 사례 수가 충분하고 모형 수렴이 가능한 조합을 우선하는 규칙에 따라 이루어졌다. 이 probe는 전체 MNLFA가 아니라 focal triangulation이며, 계수 방향과 선별 label이 모형 구조 변화에 따라 얼마나 안정적인지 확인하기 위한 보조 분석이다.",
            )
        elif txt == "표 10. 대표 후보 사례":
            set_paragraph_text(p, "표 12. focal 후보 사례")
        elif txt.startswith("사분면별 대표 사례 8개에 대해"):
            set_paragraph_text(p, txt.replace("사분면별 대표 사례", "사분면별 focal 사례"))


def style_document(doc):
    section = doc.sections[0]
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.4)
    section.right_margin = Cm(2.4)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "맑은 고딕"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    normal.paragraph_format.line_spacing = 1.18
    normal.paragraph_format.space_after = Pt(5)

    for i in range(1, 4):
        st = styles[f"Heading {i}"]
        st.font.name = "맑은 고딕"
        st._element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
        st.font.color.rgb = RGBColor(31, 55, 99)
        st.font.bold = True
        st.font.size = Pt(13 if i == 1 else 11.5)
        st.paragraph_format.space_before = Pt(12 if i == 1 else 8)
        st.paragraph_format.space_after = Pt(5)

    for idx, p in enumerate(doc.paragraphs):
        txt = p.text.strip()
        if not txt:
            continue
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
        p.paragraph_format.line_spacing = 1.18
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(5)
        p.paragraph_format.first_line_indent = None
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT

        if idx == 0:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(2)
            for r in p.runs:
                r.font.size = Pt(15.5)
                r.bold = True
                r.font.color.rgb = RGBColor(31, 55, 99)
        elif idx == 1:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(14)
            for r in p.runs:
                r.font.size = Pt(10.5)
                r.font.color.rgb = RGBColor(89, 89, 89)
        elif txt in {"국문초록", "English Abstract", "1. 서론", "2. 연구방법", "3. 결과", "4. 논의", "5. 한계", "6. 결론", "References"}:
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(6)
            for r in p.runs:
                r.font.size = Pt(13)
                r.bold = True
                r.font.color.rgb = RGBColor(31, 55, 99)
        elif re.match(r"^[1-6]\.\d+", txt):
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(4)
            for r in p.runs:
                r.font.size = Pt(11)
                r.bold = True
                r.font.color.rgb = RGBColor(31, 55, 99)
        elif txt.startswith(chr(0xD45C) + " "):  # 표
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(7)
            p.paragraph_format.space_after = Pt(3)
            for r in p.runs:
                r.font.size = Pt(9.3)
                r.bold = True
                r.font.color.rgb = RGBColor(68, 68, 68)
        elif txt.startswith("주."):
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(6)
            for r in p.runs:
                r.font.size = Pt(9)
                r.font.color.rgb = RGBColor(89, 89, 89)
        elif txt.startswith("주요어:") or txt.startswith("Keywords:"):
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for r in p.runs:
                r.font.size = Pt(9.5)
                r.font.color.rgb = RGBColor(68, 68, 68)
        else:
            for r in p.runs:
                r.font.size = Pt(10.5)
                r.bold = False
                r.font.color.rgb = RGBColor(0, 0, 0)

        for r in p.runs:
            r.font.name = "맑은 고딕"
            r._element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")

    for table in doc.tables:
        table.style = "Table Grid"
        table.autofit = True
        if table.rows:
            repeat_table_header(table.rows[0])
        for ri, row in enumerate(table.rows):
            for cell in row.cells:
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                set_cell_margins(cell, 85, 95, 85, 95)
                if ri == 0:
                    set_cell_shading(cell, "EAF0F8")
                for p in cell.paragraphs:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if len(p.text.strip()) < 28 else WD_ALIGN_PARAGRAPH.LEFT
                    p.paragraph_format.space_after = Pt(0)
                    p.paragraph_format.line_spacing = 1.05
                    for r in p.runs:
                        r.font.name = "맑은 고딕"
                        r._element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
                        r.font.size = Pt(8.2)
                        if ri == 0:
                            r.bold = True
                            r.font.color.rgb = RGBColor(31, 55, 99)

    running = "생성형 언어모형과 순서형 DIF 후보 우선순위화"
    for section in doc.sections:
        for p in section.header.paragraphs:
            p.clear()
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            r = p.add_run(running)
            r.font.name = "맑은 고딕"
            r._element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
            r.font.size = Pt(8)
            r.font.color.rgb = RGBColor(130, 130, 130)
        for p in section.footer.paragraphs:
            p.clear()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(running)
            r.font.name = "맑은 고딕"
            r._element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
            r.font.size = Pt(8)
            r.font.color.rgb = RGBColor(130, 130, 130)


def final_text_cleanup(doc):
    for p in doc.paragraphs:
        txt = p.text.strip()
        if txt.startswith("응답자 유형별 분석에서는 보호자 문항과 청소년 문항의 양상이 달랐다"):
            set_paragraph_text(
                p,
                "응답자 유형별 분석에서는 보호자 문항과 청소년 문항의 양상이 달랐다(표 7). 보호자 문항에서는 기본 지시문과 엄격 지시문 모두에서 LLM AP가 이분형 키워드 기준보다 높았다. 반면 청소년 문항에서는 LLM의 우위가 나타나지 않았고, 키워드 기준과 거의 같거나 더 낮았다. 이는 LLM의 후보 생성 능력이 전체 문항 pool에서 균질하게 작동하지 않으며, 응답자 유형과 문항 맥락에 따라 달라질 수 있음을 보여준다.",
            )
        elif txt.startswith("LLM 점수의 보정 가능성도 별도로 확인하였다"):
            set_paragraph_text(
                p,
                "LLM 점수의 보정 가능성도 별도로 확인하였다(표 8). 점수가 높은 구간에서 잠정적 선별 양성률이 대체로 높아지는 경향은 있었지만, LLM 점수 자체를 경험적 확률로 해석하기에는 차이가 컸다. 예를 들어 80-100점 구간의 양성률은 기본 지시문에서 .407, 엄격 지시문에서 .479였다. 따라서 본 연구에서 LLM 점수는 보정된 확률이 아니라 후보 우선순위화를 위한 순위 신호로 해석한다.",
            )
        elif txt.startswith("LLM의 우선순위화 성능도 제한적이었다."):
            set_paragraph_text(p, p.text.replace("이러한  결과", "이러한 결과"))


def remove_trailing_empty_paragraphs(doc):
    body = doc.element.body
    # Keep the final section properties, but remove blank paragraphs that precede it.
    children = list(body.iterchildren())
    for child in reversed(children):
        tag = child.tag.split("}")[-1]
        if tag == "sectPr":
            continue
        if tag != "p":
            break
        texts = [n.text or "" for n in child.iter() if n.tag == qn("w:t")]
        has_break = any(n.tag == qn("w:br") for n in child.iter())
        if "".join(texts).strip() == "" and not has_break:
            body.remove(child)
        else:
            break


def main():
    doc = Document(SRC)
    set_paragraph_text(doc.paragraphs[0], TITLE)
    set_paragraph_text(doc.paragraphs[1], SUBTITLE)
    apply_text_polish(doc)
    normalize_table_numbers(doc)
    final_text_cleanup(doc)
    style_document(doc)
    remove_trailing_empty_paragraphs(doc)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
