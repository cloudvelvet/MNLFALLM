from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor


SRC = r"N:\개인\다음논문\LLM_DIF_full_manuscript_draft_psychometric_style_정리본.docx"
OUT = r"N:\개인\다음논문\LLM_DIF_full_manuscript_draft_psychometric_style_정리본_rationale낮춤.docx"


def set_text(p, text):
    p.clear()
    r = p.add_run(text)
    r.font.name = "맑은 고딕"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
    r.font.size = Pt(10.5)
    return r


def style_para(p):
    txt = p.text.strip()
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    p.paragraph_format.line_spacing = 1.18
    p.paragraph_format.space_after = Pt(5)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for r in p.runs:
        r.font.name = "맑은 고딕"
        r._element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
        if txt in {"국문초록", "English Abstract", "1. 서론", "2. 연구방법", "3. 결과", "4. 논의", "5. 한계", "6. 결론", "References"}:
            r.font.size = Pt(13)
            r.bold = True
            r.font.color.rgb = RGBColor(31, 55, 99)
        elif txt.startswith(tuple(f"{i}." for i in range(1, 7))):
            r.font.size = Pt(11)
            r.bold = True
            r.font.color.rgb = RGBColor(31, 55, 99)
        elif txt.startswith("표 "):
            r.font.size = Pt(9.3)
            r.bold = True
            r.font.color.rgb = RGBColor(68, 68, 68)
        elif txt.startswith("주."):
            r.font.size = Pt(9)
            r.font.color.rgb = RGBColor(89, 89, 89)
        else:
            r.font.size = Pt(10.5)
            r.bold = False
            r.font.color.rgb = RGBColor(0, 0, 0)


def main():
    doc = Document(SRC)

    for p in doc.paragraphs:
        txt = p.text.strip()

        if txt.startswith("본 연구의 연구질문은 다음과 같다."):
            set_text(
                p,
                "본 연구의 연구질문은 다음과 같다. 첫째, LLM은 경험적 DIF 결과를 보지 않은 상태에서 청소년 및 보호자 문항의 문항-공변량 조합별 DIF 후보 가설을 구조화할 수 있는가. 둘째, LLM의 가능성 점수는 이분형 키워드 비교 기준보다 잠정적 경험 DIF 후보를 더 잘 우선순위화하는가. 셋째, 이 성능은 공변량과 응답자 유형의 특성에 따라 달라지는가. 넷째, 엄격한 DIF 구분 지시문은 기본 지시문에 비해 LLM의 우선순위화 성능과 상위 후보 목록의 안정성을 개선하는가. 판단 근거 분석은 이 연구질문에 대한 타당성 검증이 아니라, 지시문 조건에 따라 LLM 설명에서 나타나는 보조적 failure-mode 신호가 어떻게 달라지는지 기술하기 위한 부가 분석으로 다룬다.",
            )

        elif txt == "2.6 판단 근거 코딩과 활용 가능성":
            set_text(p, "2.6 LLM 판단 근거의 보조 failure-mode audit")

        elif txt.startswith("LLM 판단 근거는 별도의 분석 대상으로 코딩하였다."):
            set_text(
                p,
                "LLM 판단 근거는 DIF 원인을 입증하는 자료가 아니라, LLM이 문항-공변량 관계를 어떤 설명 양식으로 구성하는지 보여주는 텍스트 자료로 보았다. 본 연구는 사전에 정의한 자동 규칙에 따라 응답 과정 기반 설명, 검증 가능한 가설 표현, 동일 잠재특성 수준 조건의 명시, 모호한 일반론, 성별·연령 고정관념 가능성, 문항 내용과 직접 관련성이 약한 추측, impact를 DIF로 오인할 가능성을 탐색적으로 분류하였다.",
            )

        elif txt.startswith("자동 규칙 기반 코딩은 판단 근거의 설명 양식을 요약하기 위한 절차"):
            set_text(
                p,
                "이 절차는 판단 근거의 심리측정학적 타당성을 검증하기 위한 것이 아니라, LLM 설명에서 나타나는 잠재적 위험 신호를 기술하기 위한 보조 audit이다. 자동 코딩의 해석 가능성을 점검하기 위해 116개 문항-공변량-지시문 조합을 층화 표집하여 수동 검증용 파일을 마련했지만, 독립 코딩이 완료되지 않은 상태에서는 자동 코딩 결과를 판단 근거의 타당성 증거로 해석하지 않는다. 따라서 본 연구에서 판단 근거 분석은 주요 성능 평가가 아니라 failure-mode를 해석하기 위한 보조 분석으로 제한한다.",
            )

        elif txt == "3.6 판단 근거 코딩 결과":
            set_text(p, "3.6 보조 failure-mode audit 결과")

        elif txt.startswith("LLM 판단 근거에 대한 자동 코딩 결과"):
            set_text(
                p,
                "다음 결과는 LLM 판단 근거의 타당성을 검증한 것이 아니라, 자동 규칙에 따라 설명 양식과 잠재적 위험 신호의 분포를 기술한 것이다. 자동 audit 결과, 엄격한 DIF 구분 지시문은 일부 설명 양식 지표의 분포를 변화시켰다. 응답 과정 관련 근거는 기본 지시문 조건에서 89.5%였고, 엄격한 DIF 구분 지시문 조건에서 94.8%로 증가하였다. 검증 가능한 가설 표현도 68.9%에서 71.8%로 증가하였다. 반대로 같은 잠재특성 수준 조건을 명시하지 않는 판단 근거는 81.3%에서 83.2%로 약간 증가하였다.",
            )

        elif txt.startswith("그러나 엄격한 DIF 구분 지시문이 모든 위험 신호를 줄인 것은 아니다."):
            set_text(
                p,
                "그러나 이 결과를 판단 근거의 질적 개선으로 해석해서는 안 된다. 엄격한 DIF 구분 지시문이 일부 응답 과정 관련 표현을 늘렸더라도, 자동 규칙 기반 지표만으로 LLM 설명의 심리측정학적 적절성을 판단할 수는 없다. 본 결과는 strict prompt가 LLM 설명의 표면적 양식을 바꿀 수 있음을 보여주는 보조 자료이며, LLM rationale의 타당성은 별도의 전문가 검토와 후속 경험 검증을 필요로 한다.",
            )

        elif txt == "표 10. LLM 판단 근거의 자동 코딩 기반 설명 양식 요약":
            set_text(p, "표 10. LLM 판단 근거의 보조 failure-mode audit 결과")

        elif txt == "3.7 대표 후보 사례":
            set_text(p, "3.7 Focal 후보 사례")

        elif txt.startswith("LLM의 우선순위화 성능도 제한적이었다."):
            set_text(
                p,
                "LLM의 우선순위화 성능도 제한적이었다. 1-5차년도 pooled 분석에서는 LLM AP가 이분형 키워드 기준을 소폭 상회했지만, 문항 단위 bootstrap 구간은 0을 포함했고 permutation test에서도 차이가 뚜렷하지 않았다. 6차년도 단일 wave 분석에서는 키워드 기준이 LLM보다 높은 AP를 보였다. 따라서 본 연구는 LLM이 키워드 기준보다 우수하다는 결론을 지지하지 않으며, 이 자료와 이 모형 및 지시문 조건에서 LLM의 추가 신호가 불안정했음을 나타낸다. 이러한 결과는 Maeda와 Lu(2025)의 접근과 비교할 때 설계 차이의 맥락에서 이해할 수 있다. Maeda와 Lu는 경험적 DIF 레이블로 훈련된 인코더 모형을 사용해 문항 텍스트에서 DIF를 예측했으나, 본 연구의 생성형 LLM은 레이블 없이 문항-공변량 의미 관계만으로 후보를 정렬하였다. 레이블 없는 조건이 더 어려운 과제라는 점에서, 본 연구의 제한적 성능은 생성형 LLM 자체의 일반적 한계라기보다 label-free candidate prioritization이라는 과제 설정의 어려움으로 해석하는 것이 적절하다.",
            )

        elif txt.startswith("판단 근거 코딩 역시 이러한 제한적 해석 안에 놓아야 한다."):
            set_text(
                p,
                "보조 failure-mode audit 역시 이러한 제한적 해석 안에 놓아야 한다. 본 연구의 자동 audit는 LLM 설명의 타당성을 검증한 것이 아니라, 설명에서 나타나는 위험 신호와 설명 양식을 탐색적으로 요약한 절차다. LLM은 문항 표현이나 응답 과정에 근거한 설명을 만들기도 했지만, 잠재특성의 실제 차이를 문항 기능 차이처럼 서술하거나 성별·연령·소득에 관한 일반론을 문항 기능 차이로 연결할 위험도 보였다. 따라서 LLM rationale은 설득력 있는 문장이라는 이유로 증거가 될 수 없으며, 별도의 전문가 검토와 경험적 검증을 요구하는 자료로 보아야 한다.",
            )

        elif txt.startswith("셋째, 본 연구는 단일 LLM인 Gemini 2.5 Flash와 두 지시문 조건"):
            set_text(
                p,
                "셋째, 본 연구는 단일 LLM인 Gemini 2.5 Flash와 두 지시문 조건에 기반한다. 폐쇄형 LLM은 모델 버전과 API 설정이 시간에 따라 바뀔 수 있으며, MAPS 관련 문헌이나 유사한 DIF 연구가 훈련자료에 포함되었을 가능성도 배제할 수 없다. 따라서 본 연구의 blind condition은 입력 단계의 blinding을 의미하며, 훈련자료 수준의 완전한 격리를 의미하지 않는다. 또한 판단 근거 분석은 자동 규칙 기반 audit로 수행되었으므로, LLM rationale의 심리측정학적 타당성을 검증한 결과가 아니다. 전문가 맹검 코딩이 추가된다면 LLM 설명의 실제 검토 가능성과 오류 유형을 더 직접적으로 평가할 수 있다.",
            )

        elif txt.startswith("넷째, 키워드 기준과 TF-IDF 의미 유사도 기준"):
            fixed = txt if txt.endswith(".") else txt + "."
            set_text(p, fixed)

        style_para(p)

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
