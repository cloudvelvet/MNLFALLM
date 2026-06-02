import os
from copy import deepcopy

import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.shared import Inches, Pt
from docx.text.paragraph import Paragraph


SRC_NAME = "LLM_DIF_full_manuscript_draft_psychometric_style_수정본_방법결과반영_글자크기정리_SSCI보강.docx"
OUT_NAME = "LLM_DIF_full_manuscript_draft_psychometric_style_수정본_방법결과반영_글자크기정리_SSCI보강_리뷰반영.docx"
DOC_DIR = r"N:\개인\다음논문"
OUT_DIR = r"C:\chen_bauer_2024\MNLFALLM\llm_dif_output"


TITLE = "생성형 LLM은 순서형 DIF 후보를 키워드 기준보다 잘 우선순위화하는가?"
SUBTITLE = "다문화청소년패널 문항-공변량 조합의 벤치마크와 경계조건 분석"

KOR_ABSTRACT = [
    "본 연구는 다문화청소년패널(MAPS) 2기 자료의 청소년 및 보호자 문항을 대상으로, 생성형 대규모 언어모형이 순서형 DIF 검토를 위한 문항-공변량 후보를 키워드 기준보다 더 잘 우선순위화할 수 있는지 평가하였다. 분석 단위는 문항-공변량 조합이었다. Gemini 2.5 Flash에는 문항 문장, 구인 맥락, 응답 범주, 공변량 정의를 제공하되, 경험적 DIF 선별 결과, p값, 효과크기, 문항 통계량은 제공하지 않았다. LLM 산출물은 기본 지시문과 엄격한 DIF 구분 지시문 조건에서 생성되었고, 이분형 키워드 기준, TF-IDF n-gram cosine similarity 기준, 잠정적 순서형 DIF 선별 결과와 비교하였다.",
    "주 분석에서 LLM의 AP는 기본 지시문 .382, 엄격 지시문 .405였으며, 이분형 키워드 기준의 평균 AP는 각각 .375와 .384였다. 그러나 문항 단위 bootstrap 구간은 0을 포함했고, paired permutation test에서도 LLM의 우위는 통계적으로 뚜렷하지 않았다. 6차년도 단일 wave 민감도 분석에서는 오히려 키워드 기준이 LLM보다 높은 AP를 보였다. 또한 두 지시문 조건의 상위 후보 목록은 크게 달랐으며, top-5와 top-10 overlap은 모두 0이었다. LLM 점수는 TF-IDF 유사도 기준보다는 높았지만, 투명한 키워드 기준을 안정적으로 넘어서지는 못했다.",
    "이 결과는 LLM을 DIF 판정 도구로 사용하는 것을 지지하지 않는다. 본 연구의 핵심 결론은 LLM의 일반적 우수성이 아니라, LLM 기반 후보 생성의 경계조건이다. LLM은 문항-공변량 조합에 대한 검토 가능한 설명을 생성할 수 있지만, 그 설명과 점수는 지시문 조건에 민감하며 경험적 선별 기준과 안정적으로 정렬되지 않는다. 따라서 LLM 산출물은 단독 판단 근거가 아니라, 키워드 기준, 경험적 screening, 불확실성 평가, 후속 심리측정 검증과 함께 제한적으로 사용해야 할 후보 가설 자료로 해석되어야 한다.",
]

ENG_ABSTRACT = [
    "This study evaluated whether a generative large language model can prioritize item-covariate combinations for ordinal differential item functioning (DIF) review better than transparent lexical benchmarks. The analysis used youth and caregiver items from the second cohort of the Multicultural Adolescents Panel Study (MAPS). Gemini 2.5 Flash received item text, construct context, response categories, and covariate definitions, but not empirical DIF screening results, p-values, effect sizes, or item statistics. LLM outputs were generated under an original prompt and a stricter DIF-guardrail prompt, then compared with a binary keyword benchmark, a TF-IDF n-gram cosine-similarity benchmark, and provisional ordinal DIF screening labels.",
    "In the pooled Wave 1-5 analysis, LLM average precision (AP) was .382 under the original prompt and .405 under the strict prompt, whereas the binary keyword benchmark yielded mean AP values of .375 and .384, respectively. These apparent advantages were not statistically stable: item-cluster bootstrap intervals included zero, and paired permutation tests did not indicate reliable differences. In the Wave 6 single-wave sensitivity analysis, the keyword benchmark outperformed the LLM. The top-ranked LLM candidate lists were also highly prompt-sensitive, with zero overlap in the top 5 and top 10 candidates across the two prompt conditions. Although LLM scores exceeded the TF-IDF similarity baseline, they did not reliably outperform the hand-crafted lexical benchmark.",
    "The findings do not support using LLMs as DIF decision tools. Their main contribution is to identify boundary conditions for LLM-assisted DIF candidate generation. LLMs can produce structured, reviewable hypotheses about item-covariate combinations, but their rankings are prompt-sensitive and do not provide a statistically reliable signal beyond transparent lexical screening under the present design. LLM outputs should therefore be treated as provisional hypothesis material, to be interpreted alongside keyword benchmarks, empirical screening, uncertainty analyses, and subsequent psychometric validation.",
]

DISCUSSION = [
    "본 연구는 MAPS 문항-공변량 조합을 대상으로, 경험적 DIF 선별 결과를 보지 않은 LLM이 후속 검토 후보를 우선순위화할 수 있는지 평가하였다. 결과를 종합하면, 본 연구의 주된 발견은 LLM의 예측 우위가 아니라 LLM 기반 후보 생성의 경계조건이다. LLM은 구조화된 후보 설명을 생성할 수 있었지만, 투명한 어휘 기준을 안정적으로 넘어서는 독립적 신호를 제공하지는 못했다.",
    "가장 중요한 결과는 지시문 민감도였다. 기본 지시문과 엄격한 DIF 구분 지시문 사이의 Spearman 순위상관은 중간 수준이었지만, 상위 후보 목록은 거의 겹치지 않았다. 특히 top-5와 top-10 overlap은 모두 0이었다. 이는 LLM이 제안하는 실제 검토 목록이 지시문 조건에 크게 의존한다는 뜻이다. DIF source investigation의 목적이 제한된 전문가 검토 자원을 상위 후보에 배분하는 것이라면, 이러한 상위 목록의 불안정성은 단순한 기술적 문제가 아니라 실무적 한계다.",
    "LLM의 우선순위화 성능도 제한적이었다. 1-5차년도 pooled 분석에서는 LLM AP가 이분형 키워드 기준을 소폭 상회했지만, 문항 단위 bootstrap 구간은 0을 포함했고 permutation test에서도 차이가 뚜렷하지 않았다. 6차년도 단일 wave 분석에서는 키워드 기준이 LLM보다 높은 AP를 보였다. 따라서 본 연구는 LLM이 키워드 기준보다 우수하다는 결론을 지지하지 않는다. 더 정확한 결론은, 이 자료와 이 모형 및 지시문 조건에서 LLM의 추가 신호가 불안정했다는 것이다.",
    "LLM 점수의 보정 양상도 신중하게 해석해야 한다. 높은 LLM 점수 구간에서 잠정적 선별 양성률이 대체로 높아지는 경향은 있었지만, 점수 자체가 경험적 확률로 해석될 정도로 보정되어 있지는 않았다. 예를 들어 80-100점 구간의 양성률은 기본 지시문에서 .407, 엄격 지시문에서 .479였다. 이는 LLM의 80점대 점수가 실제 80% 수준의 경험적 양성 가능성을 뜻하지 않음을 보여준다. LLM 점수는 확률 추정치라기보다 순위 신호로 제한해 해석하는 편이 적절하다.",
    "응답자 유형별 결과는 LLM의 역할이 문항 맥락에 따라 달라질 수 있음을 보여준다. 보호자 문항에서는 LLM AP가 키워드 기준보다 높았지만, 청소년 문항에서는 LLM의 우위가 나타나지 않았다. 보호자 문항에서는 한국어 능력, 문화적응, 가족 맥락이 문항 해석과 더 직접적으로 연결될 가능성이 있다. 반면 청소년 문항에서는 또래 관계, 학교생활, 발달 단계가 얽혀 있어 LLM이 생성한 일반적 설명이 실제 선별 신호와 안정적으로 맞물리지 않았을 수 있다. 이 차이는 후속 연구에서 응답자 유형과 구인을 분리해 검토할 필요를 보여준다.",
    "사분면 분석은 실무적 사용 조건을 더 분명하게 보여준다. LLM과 키워드 기준이 모두 높은 both-high 영역은 경험적 양성률이 가장 높았고, both-low 영역은 가장 낮았다. 반면 LLM-high/keyword-low 영역은 both-high만큼 강한 후보군을 만들지 못했다. keyword-high/LLM-low 영역에서도 양성 사례가 적지 않았다. 따라서 LLM 점수만으로 키워드 후보를 배제하거나, LLM 단독 고득점 후보를 우선 검토 대상으로 삼는 전략은 조심해야 한다. 보다 방어 가능한 절차는 키워드 기준으로 투명한 1차 후보를 만들고, LLM을 이용해 그 후보의 설명 가능성과 우선순위를 보조적으로 검토하는 방식이다.",
    "판단 근거 코딩 역시 이러한 제한적 해석 안에 놓아야 한다. 본 연구의 자동 코딩은 LLM 설명의 타당성을 검증한 것이 아니라, 설명에서 나타나는 위험 신호와 설명 양식을 탐색적으로 요약한 audit 절차다. LLM은 문항 표현이나 응답 과정에 근거한 설명을 만들기도 했지만, 잠재특성의 실제 차이를 문항 기능 차이처럼 서술하거나 성별·연령·소득에 관한 일반론을 문항 기능 차이로 연결할 위험도 보였다. 따라서 LLM rationale은 설득력 있는 문장이라는 이유로 증거가 될 수 없으며, 별도의 전문가 검토와 경험적 검증을 요구하는 자료로 보아야 한다.",
    "결국 본 연구의 기여는 LLM이 DIF 후보를 잘 찾아낸다는 주장이 아니다. 오히려 LLM 산출물을 후보 가설로 제한하고, 키워드 기준, TF-IDF 유사도 기준, 지시문 민감도, bootstrap과 permutation, 6차년도 민감도 분석, 사분면 분석, focal random-intercept probe를 함께 사용해 LLM의 경계조건을 드러낸 데 있다. 이는 LLM을 도입하려는 연구자에게 긍정적 사용법만이 아니라 사용을 중단하거나 보류해야 할 조건까지 제시한다는 점에서 의미가 있다.",
]

LIMITATIONS = [
    "본 연구의 결과는 다음 제한 안에서 해석되어야 한다. 첫째, 잠정적 순서형 DIF 선별 양성 조합은 DIF의 참값이 아니라 경험적 검토 후보이다. 선별 기준은 proportional-odds 누적로짓 모형, leave-one-item-out 척도 점수 proxy, FDR 기준, practical nonzero 기준에 의존한다. 본 연구는 cluster-robust 분석, 6차년도 단일 wave 분석, bootstrap, permutation, focal random-intercept probe를 통해 민감도를 점검했지만, 이 절차가 완전한 psychometric validation을 대체하지는 않는다.",
    "둘째, proportional-odds 가정은 문항별로 별도 검정하지 않았다. 특정 공변량의 효과가 응답 범주 경계마다 다르게 나타나는 경우, 본 연구의 screening label은 일부 조합을 과소 또는 과대 선별할 수 있다. 따라서 본 연구의 경험적 기준은 threshold-specific DIF의 확정 검정이 아니라, uniform ordinal DIF 후보를 탐색하기 위한 잠정적 기준으로 이해해야 한다.",
    "셋째, leave-one-item-out 척도 점수 proxy는 잠재특성의 직접 추정치가 아니다. 척도 구조가 다차원적이거나 문항 수가 적은 경우 matching variable의 측정오차가 커질 수 있으며, DIF 가능성이 있는 다른 문항이 proxy에 포함되면 anchor contamination이 발생할 수 있다. 후속 연구에서는 IRT 기반 theta, anchor purification, MNLFA 또는 random-intercept ordinal model을 사용해 기준 label을 더 엄격하게 구성할 필요가 있다.",
    "넷째, 본 연구는 단일 LLM인 Gemini 2.5 Flash와 두 지시문 조건에 기반한다. 폐쇄형 LLM은 모델 버전과 API 설정이 시간에 따라 바뀔 수 있으며, 연구자가 훈련자료 포함 여부를 완전히 확인할 수 없다. MAPS 관련 문헌이나 유사한 DIF 연구가 훈련자료에 포함되었을 가능성도 배제할 수 없다. 따라서 본 연구의 blind condition은 입력 단계의 blinding을 의미하며, 훈련자료 수준의 완전한 격리를 의미하지 않는다.",
    "다섯째, 판단 근거 코딩은 자동 규칙 기반으로 수행되었다. 자동 코딩은 LLM 설명의 양식과 위험 신호를 요약하는 audit 절차이지, 판단 근거의 심리측정학적 타당성을 입증하는 절차가 아니다. 전문가 맹검 코딩이 추가된다면 LLM rationale의 실제 검토 가능성과 오류 유형을 더 직접적으로 평가할 수 있다.",
    "여섯째, 키워드 기준과 TF-IDF 의미 유사도 기준은 모두 투명한 비교 기준이지만 중립적 기준은 아니다. 키워드 목록은 연구자의 판단으로 구성되었고, TF-IDF 유사도는 문맥적 추론을 충분히 반영하지 못한다. 특히 차별 경험이나 한국어 능력처럼 표면 어휘가 강한 공변량에서는 키워드 기준이 강한 benchmark로 작동할 수 있다. 이 점은 LLM의 약점만이 아니라, DIF 후보 선별에서 단순하고 투명한 기준이 갖는 실질적 가치를 보여주는 결과이기도 하다.",
]

CONCLUSION = [
    "본 연구는 LLM을 DIF 판정 도구로 지지하지 않는다. LLM은 일부 조건에서 구조화된 후보 설명을 생성하고 TF-IDF 유사도 기준보다 높은 순위 성능을 보였지만, 투명한 키워드 기준을 안정적으로 능가하지 못했고 상위 후보 목록은 지시문에 크게 흔들렸다. 따라서 LLM 점수는 경험적 DIF의 대체물이 아니라 검토 후보를 정리하기 위한 보조 신호로 제한되어야 한다.",
    "본 연구의 결론은 보수적이지만 실무적으로 분명하다. 대규모 문항-공변량 공간에서 LLM을 사용하려면, 첫째 투명한 키워드 또는 어휘 기준을 반드시 함께 제시해야 하고, 둘째 LLM 단독 고득점보다 LLM과 baseline이 수렴하는 후보를 우선 검토해야 하며, 셋째 지시문 변화에 따른 상위 후보 목록의 안정성을 보고해야 한다. 이러한 조건을 충족하지 않는 LLM 기반 DIF 후보 목록은 연구자의 검토를 돕기보다 검토 우선순위를 임의로 바꿀 위험이 있다.",
    "따라서 본 연구의 기여는 LLM의 성능을 옹호하는 데 있지 않다. 기여는 LLM을 심리측정 검증 앞단에 배치할 때 필요한 비교 기준과 안전장치를 경험적으로 제시한 데 있다. LLM은 답을 내리는 도구가 아니라, 검증해야 할 설명을 생산하는 도구다. 그 설명의 가치는 투명한 baseline, 경험적 screening, 불확실성 평가, 후속 심리측정 모형 속에서만 판단될 수 있다.",
]


def set_paragraph_text(paragraph, text):
    paragraph.clear()
    run = paragraph.add_run(text)
    run.font.name = "맑은 고딕"
    run.font.size = Pt(10.5)


def paragraph_after(paragraph, text="", style=None):
    new_p = deepcopy(paragraph._p)
    paragraph._p.addnext(new_p)
    p = paragraph._parent.paragraphs[[x._p for x in paragraph._parent.paragraphs].index(new_p)]
    p.clear()
    if style:
        p.style = style
    run = p.add_run(text)
    run.font.name = "맑은 고딕"
    run.font.size = Pt(10.5)
    return p


def insert_paragraph_before(paragraph, text, style=None, size=10.5, bold=False):
    new_p = deepcopy(paragraph._p)
    paragraph._p.addprevious(new_p)
    p = paragraph._parent.paragraphs[[x._p for x in paragraph._parent.paragraphs].index(new_p)]
    p.clear()
    if style:
        p.style = style
    run = p.add_run(text)
    run.font.name = "맑은 고딕"
    run.font.size = Pt(size)
    run.bold = bold
    return p


def insert_paragraph_after_element(element, parent, text, style=None, size=10.5, bold=False):
    new_p = OxmlElement("w:p")
    element.addnext(new_p)
    p = Paragraph(new_p, parent)
    if style:
        p.style = style
    run = p.add_run(text)
    run.font.name = "맑은 고딕"
    run.font.size = Pt(size)
    run.bold = bold
    return p


def insert_table_after(paragraph, rows, headers):
    tbl = paragraph._parent.add_table(rows=1, cols=len(headers), width=Inches(6.5))
    tbl.style = "Table Grid"
    for j, h in enumerate(headers):
        cell = tbl.rows[0].cells[j]
        cell.text = h
        for p in cell.paragraphs:
            for r in p.runs:
                r.font.name = "맑은 고딕"
                r.font.size = Pt(8.5)
                r.bold = True
    for row in rows:
        cells = tbl.add_row().cells
        for j, h in enumerate(headers):
            cells[j].text = str(row[h])
            for p in cells[j].paragraphs:
                for r in p.runs:
                    r.font.name = "맑은 고딕"
                    r.font.size = Pt(8.5)
    paragraph._p.addnext(tbl._tbl)
    return tbl


def replace_section(doc, start_heading, end_heading, new_paragraphs):
    paras = doc.paragraphs
    start = next(i for i, p in enumerate(paras) if p.text.strip() == start_heading)
    end = next(i for i, p in enumerate(paras[start + 1 :], start + 1) if p.text.strip() == end_heading)
    set_paragraph_text(paras[start], start_heading)
    paras[start].runs[0].font.size = Pt(13)
    paras[start].runs[0].bold = True
    first = paras[start + 1]
    set_paragraph_text(first, new_paragraphs[0])
    for p in paras[start + 2 : end]:
        p._element.getparent().remove(p._element)
    anchor = first
    for text in new_paragraphs[1:]:
        anchor = paragraph_after(anchor, text)


def main():
    src = os.path.join(DOC_DIR, SRC_NAME)
    out = os.path.join(DOC_DIR, OUT_NAME)
    doc = Document(src)

    set_paragraph_text(doc.paragraphs[0], TITLE)
    doc.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.paragraphs[0].runs[0].font.size = Pt(16)
    doc.paragraphs[0].runs[0].bold = True
    set_paragraph_text(doc.paragraphs[1], SUBTITLE)
    doc.paragraphs[1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.paragraphs[1].runs[0].font.size = Pt(11)

    # Korean abstract: heading at 3, content 4-6, keywords 7.
    for idx, text in zip([4, 5, 6], KOR_ABSTRACT):
        set_paragraph_text(doc.paragraphs[idx], text)
    set_paragraph_text(
        doc.paragraphs[7],
        "주요어: 차별기능문항, 대규모 언어모형, 다문화청소년패널, 키워드 기준, 지시문 민감도, 심리측정 벤치마크",
    )

    for idx, text in zip([9, 10, 11], ENG_ABSTRACT):
        set_paragraph_text(doc.paragraphs[idx], text)
    set_paragraph_text(
        doc.paragraphs[12],
        "Keywords: differential item functioning, large language models, multicultural adolescents, lexical benchmark, prompt sensitivity, psychometric benchmarking",
    )

    # Strengthen introduction framing by replacing final purpose paragraph.
    for p in doc.paragraphs:
        if p.text.strip().startswith("본 연구는 MAPS의 청소년 및 보호자 응답 문항을 대상으로"):
            set_paragraph_text(
                p,
                "본 연구는 MAPS의 청소년 및 보호자 응답 문항을 대상으로, 경험적 DIF 결과를 모르는 LLM이 문항-공변량 조합별 후보를 어떻게 우선순위화하는지 평가한다. 분석의 초점은 LLM이 DIF를 판정할 수 있는지가 아니라, LLM 점수가 투명한 어휘 기준을 넘어서는 안정적인 검토 신호를 제공하는지, 그리고 지시문 변화에 따라 후보 목록이 얼마나 달라지는지에 있다. 따라서 본 연구는 LLM의 가능성을 주장하기보다, 심리측정 검증 앞단에서 LLM을 사용할 때 필요한 비교 기준과 경계조건을 밝히는 데 목적을 둔다.",
            )
            break

    # Insert respondent-type and calibration summaries before old 3.4.
    cal = pd.read_csv(os.path.join(OUT_DIR, "maps_llm_score_calibration_bins.csv"))
    resp = pd.read_csv(os.path.join(OUT_DIR, "maps_llm_respondent_type_ap.csv"))
    old_34 = next(p for p in doc.paragraphs if p.text.strip() == "3.4 지시문 민감도")
    h = insert_paragraph_before(old_34, "3.4 응답자 유형별 성능과 LLM 점수 보정", size=11, bold=True)
    p1 = paragraph_after(
        h,
        "응답자 유형별 분석에서는 보호자 문항과 청소년 문항의 양상이 달랐다(표 6-1). 보호자 문항에서는 기본 지시문과 엄격 지시문 모두에서 LLM AP가 이분형 키워드 기준보다 높았다. 반면 청소년 문항에서는 LLM의 우위가 나타나지 않았고, 키워드 기준과 거의 같거나 더 낮았다. 이는 LLM의 후보 생성 능력이 전체 문항 pool에서 균질하게 작동하지 않으며, 응답자 유형과 문항 맥락에 따라 달라질 수 있음을 보여준다.",
    )
    cap1 = paragraph_after(p1, "표 6-1. 응답자 유형별 LLM과 키워드 기준의 AP 비교")
    cap1.runs[0].font.size = Pt(9.5)
    tbl1 = insert_table_after(cap1, resp.to_dict("records"), list(resp.columns))
    p2 = insert_paragraph_after_element(
        tbl1._tbl,
        doc._body,
        "LLM 점수의 보정 가능성도 별도로 확인하였다(표 6-2). 점수가 높은 구간에서 잠정적 선별 양성률이 대체로 높아지는 경향은 있었지만, LLM 점수 자체를 경험적 확률로 해석하기에는 차이가 컸다. 예를 들어 80-100점 구간의 양성률은 기본 지시문에서 .407, 엄격 지시문에서 .479였다. 따라서 본 연구에서 LLM 점수는 보정된 확률이 아니라 후보 우선순위화를 위한 순위 신호로 해석한다.",
    )
    cap2 = paragraph_after(p2, "표 6-2. LLM 점수 구간별 잠정적 선별 양성률")
    cap2.runs[0].font.size = Pt(9.5)
    insert_table_after(cap2, cal.to_dict("records"), list(cal.columns))

    heading_map = {
        "3.4 지시문 민감도": "3.5 지시문 민감도",
        "3.5 판단 근거 코딩 결과": "3.6 판단 근거 코딩 결과",
        "3.6 대표 후보 사례": "3.7 대표 후보 사례",
        "3.7 Screening 기준 민감도 분석": "3.8 Screening 기준 민감도 분석",
        "3.8 6차년도 단일 wave 민감도 분석": "3.9 6차년도 단일 wave 민감도 분석",
        "3.9 불확실성 평가": "3.10 불확실성 평가",
        "3.10 의미 유사도 기준과의 비교": "3.11 의미 유사도 기준과의 비교",
        "3.11 LLM-키워드 사분면과 focal random-intercept probe": "3.12 LLM-키워드 사분면과 focal random-intercept probe",
    }
    for p in doc.paragraphs:
        txt = p.text.strip()
        if txt in heading_map:
            set_paragraph_text(p, heading_map[txt])
            p.runs[0].font.size = Pt(11)
            p.runs[0].bold = True

    # Reframe discussion, limitations, conclusion.
    replace_section(doc, "4. 논의", "5. 한계", DISCUSSION)
    replace_section(doc, "5. 한계", "6. 결론", LIMITATIONS)
    replace_section(doc, "6. 결론", "References", CONCLUSION)

    # General font cleanup.
    for p in doc.paragraphs:
        txt = p.text.strip()
        if not txt:
            continue
        for r in p.runs:
            r.font.name = "맑은 고딕"
            if txt == TITLE:
                r.font.size = Pt(16)
                r.bold = True
            elif txt == SUBTITLE:
                r.font.size = Pt(11)
            elif txt in {"국문초록", "English Abstract", "1. 서론", "2. 연구방법", "3. 결과", "4. 논의", "5. 한계", "6. 결론", "References"}:
                r.font.size = Pt(13)
                r.bold = True
            elif txt.startswith("표 ") or txt.startswith("주."):
                r.font.size = Pt(9.5)
            else:
                r.font.size = Pt(10.5)

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    for r in p.runs:
                        r.font.name = "맑은 고딕"
                        r.font.size = Pt(8.5)

    doc.save(out)
    print(out)


if __name__ == "__main__":
    main()
