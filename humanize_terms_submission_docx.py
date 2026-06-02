from pathlib import Path
import shutil

from docx import Document


WORK = Path(r"C:\chen_bauer_2024\docx_work")
SRC = WORK / "submission_trim_tables_aligned.docx"
LOCAL_OUT = WORK / "submission_trim_humanized_terms.docx"
FINAL_OUT = Path(r"N:\개인\다음논문\LLM_DIF_full_manuscript_draft_psychometric_style_Gemini범위수정_투고압축본_최종.docx")


REPLACEMENTS = [
    ("JSON schema", "JSON 출력 형식"),
    ("JSON 출력 schema", "JSON 출력 형식"),
    ("naive baseline", "단순 기준"),
    ("hand-crafted lexical benchmark", "연구자 구성 어휘 기준"),
    ("lexical benchmark", "어휘 기준"),
    ("benchmark", "비교 기준"),
    ("strict prompt", "엄격 지시문"),
    ("base prompt", "기본 지시문"),
    ("prompt engineering", "지시문 조정"),
    ("prompt sensitivity", "지시문 민감도"),
    ("Keyword-high", "어휘 기준 high"),
    ("LLM-high/lexical-low", "LLM-high/어휘-low"),
    ("lexical-high/LLM-low", "어휘-high/LLM-low"),
    ("random-intercept ordinal probe", "무선절편 순서형 보조 분석"),
    ("random-intercept probe", "무선절편 보조 분석"),
    ("이 probe는", "이 보조 분석은"),
    ("probe 결과", "보조 분석 결과"),
    ("probe는", "보조 분석은"),
    ("이 probe", "이 보조 분석"),
    ("high-risk 후보", "고위험 후보"),
    ("최적 분류 threshold", "최적 분류 절단점"),
    ("screening", "선별"),
    ("parsing", "파싱"),
    ("raw response", "원자료 응답"),
]


def replace_in_paragraph(paragraph):
    if not paragraph.runs:
        return
    full = "".join(run.text for run in paragraph.runs)
    new = full
    for old, repl in REPLACEMENTS:
        new = new.replace(old, repl)
    if new == full:
        return

    # Keep the paragraph-level style. Most edited paragraphs in this manuscript
    # are plain body/caption paragraphs, so rebuilding the runs is acceptable.
    first = paragraph.runs[0]
    for run in paragraph.runs[1:]:
        run.text = ""
    first.text = new


doc = Document(SRC)

for paragraph in doc.paragraphs:
    replace_in_paragraph(paragraph)

for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                replace_in_paragraph(paragraph)

doc.save(LOCAL_OUT)
shutil.copy2(LOCAL_OUT, FINAL_OUT)
print(LOCAL_OUT)
print(FINAL_OUT)
