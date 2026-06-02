from pathlib import Path
import re
import shutil

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


WORK = Path(r"C:\chen_bauer_2024\docx_work")
SRC = WORK / "submission_trim_publication_ready_black.docx"
LOCAL_OUT = WORK / "submission_trim_tables_aligned.docx"
FINAL_OUT = Path(r"N:\개인\다음논문\LLM_DIF_full_manuscript_draft_psychometric_style_Gemini범위수정_투고압축본_최종.docx")


WIDTHS = {
    1: [1.55, 5.20],
    2: [1.15, 1.70, 0.75, 1.75, 1.40],
    3: [1.75, 4.95],
    4: [1.05, 3.00, 1.25, 1.45],
    5: [0.75, 0.72, 0.72, 0.60, 0.57, 0.70, 0.60, 0.72, 0.60, 0.72],
    6: [1.25, 0.92, 0.92, 0.92, 0.92, 0.70, 0.86],
    7: [0.82, 0.82, 0.86, 0.86, 0.82, 0.76, 0.86, 0.75],
    8: [1.70, 1.70, 1.70, 1.65],
    9: [1.55, 2.65, 2.85],
}

NUMERIC = re.compile(r"^[+-]?(?:\d+|\.\d+|\d+\.\d+)(?:/\d+)?$|^-\.\d+$")


def twips(inches: float) -> str:
    return str(int(inches * 1440))


def set_tbl_width(table, width_inches=6.75):
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), twips(width_inches))
    tbl_w.set(qn("w:type"), "dxa")

    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")


def set_cell_width(cell, width_inches):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), twips(width_inches))
    tc_w.set(qn("w:type"), "dxa")


def set_cell_margins(cell, top=70, start=85, bottom=70, end=85):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def shade_cell(cell, fill="EAF0F7"):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_borders(table, color="666666", size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        node = borders.find(tag)
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = tr_pr.find(qn("w:tblHeader"))
    if tbl_header is None:
        tbl_header = OxmlElement("w:tblHeader")
        tr_pr.append(tbl_header)
    tbl_header.set(qn("w:val"), "true")


def is_number_like(text):
    t = text.strip().replace(",", "")
    return bool(NUMERIC.match(t))


def format_cell(cell, table_idx, row_idx, col_idx, total_cols):
    text = cell.text.strip()
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER if row_idx == 0 else WD_CELL_VERTICAL_ALIGNMENT.TOP
    set_cell_margins(cell)
    if row_idx == 0:
        shade_cell(cell)

    for p in cell.paragraphs:
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.05

        if row_idx == 0:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif is_number_like(text):
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif total_cols >= 7 and col_idx >= 2:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif table_idx in (5, 6, 7, 8) and col_idx > 0:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if len(text) <= 18 else WD_ALIGN_PARAGRAPH.LEFT
        else:
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT

        for run in p.runs:
            run.font.name = "맑은 고딕"
            run.font.color.rgb = RGBColor(0, 0, 0)
            run.font.size = Pt(8.8 if total_cols >= 8 else 9.2)
            if row_idx == 0:
                run.bold = True


doc = Document(SRC)

for idx, table in enumerate(doc.tables, start=1):
    widths = WIDTHS.get(idx)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_tbl_width(table, sum(widths) if widths else 6.75)
    set_borders(table)
    if table.rows:
        repeat_header(table.rows[0])

    for r, row in enumerate(table.rows):
        for c, cell in enumerate(row.cells):
            if widths and c < len(widths):
                set_cell_width(cell, widths[c])
            format_cell(cell, idx, r, c, len(row.cells))

doc.save(LOCAL_OUT)
shutil.copy2(LOCAL_OUT, FINAL_OUT)
print(LOCAL_OUT)
print(FINAL_OUT)
