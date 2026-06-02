from pathlib import Path
from docx import Document


def set_paragraph(paragraph, text):
    paragraph.clear()
    paragraph.add_run(text)


def main():
    src = list(Path("N:/").glob("개인/다음논문/LLM_DIF_full_manuscript_draft_psychometric_style_수정본.docx"))[0]
    dst = src.with_name(src.stem + "_방법결과반영.docx")
    doc = Document(src)

    set_paragraph(
        doc.paragraphs[29],
        "본 연구는 MAPS 2기 1-5차년도 자료를 주 분석 자료로 사용하였다. 자료는 청소년과 보호자를 반복 추적한 종단 패널이지만, 본 연구의 목적은 종단 측정동일성 또는 MNLFA를 직접 추정하는 데 있지 않다. 분석의 초점은 경험적 선별 결과를 LLM에 제공하지 않은 상태에서, 후속 검토의 우선순위가 될 문항-공변량 후보를 얼마나 잘 정렬할 수 있는지에 있다. 따라서 1-5차년도 자료는 문항 응답과 공변량 정보를 결합하기 위해 pooled long format으로 구성하였고, 6차년도 자료는 단일 wave 민감도 분석에 사용하였다.",
    )

    set_paragraph(
        doc.paragraphs[37],
        "LLM에는 Gemini 2.5 Flash를 사용하였다. API 호출에서는 temperature를 0으로 설정하고 응답 형식은 JSON으로 제한하였다. 입력에는 문항 ID, 응답자 유형, 척도/구인 정보, 문항 문장, 응답 범주, 공변량 정의를 포함하였다. 경험적 선별 결과, p값, 효과크기, 문항 통계량, 키워드 점수는 입력하지 않았다. 따라서 LLM 산출물은 경험적 선별 결과를 보지 않은 상태에서 생성된 후보 설명으로 간주하였다. 모든 raw response와 파싱 결과는 재현성 점검을 위해 보존하였다.",
    )

    # Table 3: raw response disclosure row.
    table3 = doc.tables[2]
    for row in table3.rows:
        if row.cells[0].text.strip() == "raw response 공개":
            row.cells[1].text = (
                "모든 raw response는 재현성 점검을 위해 보존하였다. 다만 MAPS 문항 원문 공개가 제한되는 경우에는 "
                "문항 ID, 예측값, 판단 근거, 파싱 결과를 분리하여 공개하고, 문항 원문 전체는 자료 이용 조건과 저작권 조건을 따른다."
            )
            break

    set_paragraph(
        doc.paragraphs[98],
        "잠정적 경험 DIF 선별 기준을 달리했을 때도 결론의 방향은 크게 달라지지 않았다(표 11). 기준을 완화하거나 강화하면 양성 조합 수와 양성 비율은 변하지만, LLM과 키워드 기준의 상대적 성능은 조건에 따라 달라졌다. 특히 효과크기 기준을 강하게 적용할수록 양성 조합 수가 줄어들기 때문에 AP의 안정성도 함께 낮아진다. 따라서 본 연구의 결과는 특정 임계값 하나에 의존한 확정적 결론이라기보다, 여러 잠정적 기준 아래에서 LLM의 제한적이고 조건부적인 우선순위화 가능성을 확인한 결과로 해석해야 한다.",
    )

    set_paragraph(
        doc.paragraphs[100],
        "주. 이분형 키워드 AP의 대괄호 안 값은 표본추출 불확실성에 대한 신뢰구간이 아니라, 키워드 점수의 동점 처리를 무작위로 2,000회 반복했을 때의 2.5-97.5 백분위 범위이다.",
    )

    set_paragraph(
        doc.paragraphs[101],
        "반복측정 자료의 응답자 내 의존성을 고려하기 위해 cluster-robust 표준오차를 사용한 선별 결과도 민감도 분석으로 검토하였다(표 12). 이 분석에서도 LLM과 이분형 키워드 기준의 차이는 크지 않았다. 다만 이 분석은 반복측정 의존성을 완전히 모형화한 것이 아니라, 표준오차 처리 방식 변화에 따라 screening label이 얼마나 민감하게 달라지는지를 점검하기 위한 보조 분석으로 보아야 한다.",
    )

    set_paragraph(
        doc.paragraphs[106],
        "6차년도 단일 wave만 사용한 민감도 분석에서는 주 분석과 다른 양상이 나타났다(표 13). 기본 지시문 조건에서 LLM AP는 .208, 이분형 키워드 기준의 평균 AP는 .248이었다. 엄격 지시문 조건에서도 LLM AP는 .201, 키워드 기준의 평균 AP는 .254였다. 즉 6차년도 단일 wave에서는 전체 AP 기준으로 키워드 기준이 LLM보다 높았다.",
    )

    doc.save(dst)
    print(dst)


if __name__ == "__main__":
    main()
