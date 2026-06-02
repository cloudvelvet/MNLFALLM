from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK


IN = Path(r"C:\chen_bauer_2024\docx_work\ssci_reframed.docx")
OUT = Path(r"C:\chen_bauer_2024\docx_work\ssci_reframed_fixed.docx")


def set_cell(cell, text: str) -> None:
    cell.text = text


def delete_paragraph(paragraph) -> None:
    p = paragraph._element
    p.getparent().remove(p)
    paragraph._p = paragraph._element = None


def add_after(paragraph, text: str, style=None):
    from docx.oxml import OxmlElement
    from docx.text.paragraph import Paragraph

    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    p = Paragraph(new_p, paragraph._parent)
    if style is not None:
        p.style = style
    p.add_run(text)
    return p


doc = Document(IN)

# Repair tables whose Korean labels were corrupted in an earlier edit.
t6 = doc.tables[6]
t6_data = [
    ["지시문 조건", "응답자 유형", "분석 조합 수", "선별 양성 수", "양성 비율", "LLM AP", "전문가 어휘 AP", "AP 차이"],
    ["기본", "보호자", "152", "41", ".270", ".464", ".386", ".078"],
    ["기본", "청소년", "334", "87", ".260", ".362", ".381", "-.020"],
    ["엄격", "보호자", "152", "41", ".270", ".510", ".386", ".124"],
    ["엄격", "청소년", "330", "86", ".261", ".382", ".392", "-.010"],
]
for r, row in enumerate(t6_data):
    for c, val in enumerate(row):
        set_cell(t6.cell(r, c), val)

t7 = doc.tables[7]
t7_data = [
    ["지시문 조건", "점수 구간", "분석 조합 수", "평균 LLM 점수", "선별 양성률"],
    ["기본", "0-19", "145", "10.3", ".193"],
    ["기본", "20-39", "44", "23.9", ".114"],
    ["기본", "40-59", "59", "48.8", ".237"],
    ["기본", "60-79", "179", "65.3", ".318"],
    ["기본", "80-100", "59", "83.8", ".407"],
    ["엄격", "0-19", "166", "10.0", ".175"],
    ["엄격", "20-39", "18", "20.8", ".056"],
    ["엄격", "40-59", "18", "51.4", ".278"],
    ["엄격", "60-79", "163", "66.8", ".221"],
    ["엄격", "80-100", "117", "84.6", ".479"],
]
for r, row in enumerate(t7_data):
    for c, val in enumerate(row):
        set_cell(t7.cell(r, c), val)

# Table text polish.
for tbl in doc.tables:
    for row in tbl.rows:
        for cell in row.cells:
            text = cell.text
            text = text.replace("이분형 키워드 AP", "전문가 어휘 AP")
            text = text.replace("키워드 AP", "전문가 어휘 AP")
            text = text.replace("키워드 P@5", "전문가 어휘 P@5")
            text = text.replace("키워드 P@10", "전문가 어휘 P@10")
            text = text.replace("키워드", "전문가 어휘")
            text = text.replace("keyword-low", "lexical-low")
            text = text.replace("keyword-high", "lexical-high")
            if text != cell.text:
                cell.text = text

# Paragraph text polish for remaining lexical labels.
for p in doc.paragraphs:
    text = p.text
    new = (
        text.replace("단순 키워드 매칭", "단순 어휘 매칭")
        .replace("키워드 매칭", "어휘 매칭")
        .replace("3.12 LLM-키워드", "3.12 LLM-전문가 어휘")
        .replace("키워드", "전문가 어휘")
        .replace("keyword-low", "lexical-low")
        .replace("keyword-high", "lexical-high")
    )
    if new != text:
        for run in list(p.runs):
            run.text = ""
        if p.runs:
            p.runs[0].text = new
        else:
            p.add_run(new)

# Rebuild appendices in the correct order.
for p in list(doc.paragraphs):
    if p.text.strip().startswith("부록 A") or p.text.strip().startswith("부록 B") or p.text.strip().startswith("기본 지시문은") or p.text.strip().startswith("LLM 출력은 다음 필드를"):
        delete_paragraph(p)

last = doc.paragraphs[-1]
h1 = add_after(last, "부록 A. LLM 지시문 요약", style=doc.styles["Heading 1"])
h1.runs[0].add_break(WD_BREAK.PAGE)
a_text = add_after(
    h1,
    "기본 지시문은 문항 문장, 구인 설명, 응답 범주, 공변량 정의를 함께 제시한 뒤, 해당 문항-공변량 조합에서 순서형 DIF 후보 가능성, 예상 방향, 확신도, 판단 근거를 JSON 형식으로 산출하도록 요구하였다. 엄격한 DIF 구분 지시문은 여기에 실제 잠재특성 차이(impact)와 같은 잠재특성 수준에서의 문항 기능 차이(DIF)를 구분하라는 조건을 추가하였다. 특히 단순한 집단 평균 차이, 일반적 위험요인, 사회적 배경 차이를 DIF의 근거로 사용하지 말고, 문항 표현이나 응답 과정이 응답 범주 사용의 차이로 이어질 수 있을 때만 높은 점수를 부여하도록 지시하였다.",
    style=doc.styles["Normal"],
)
h2 = add_after(a_text, "부록 B. JSON 출력 schema", style=doc.styles["Heading 1"])
add_after(
    h2,
    "LLM 출력은 다음 필드를 포함하도록 제한하였다: custom_id, model, scale_id, item_id, covariate, threshold_dif_probability_0_100, expected_direction, confidence_0_100, rationale, parse_status. Probability와 confidence는 0-100 범위의 정수로 기록하였고, expected_direction은 positive, negative, unclear 중 하나로 제한하였다. Parsing 실패, 누락 필드, 범위 밖 점수는 별도로 표시하고 분석에서 제외하거나 보정 규칙에 따라 처리하였다.",
    style=doc.styles["Normal"],
)

doc.save(OUT)
print(OUT)
