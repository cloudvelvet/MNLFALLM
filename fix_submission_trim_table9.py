from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, Inches
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


WORK = Path(r"C:\chen_bauer_2024\docx_work")
SRC = WORK / "submission_trim_final.docx"
OUT = WORK / "submission_trim_final2.docx"


def set_cell_width(cell, width_inches):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width_inches * 1440)))
    tc_w.set(qn("w:type"), "dxa")


def set_cell_margins(cell, top=80, start=90, bottom=80, end=90):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


doc = Document(SRC)

table = doc.tables[-1]
table.alignment = WD_TABLE_ALIGNMENT.LEFT
table.autofit = False

rows = [
    ["분석", "핵심 결과", "해석"],
    [
        "불확실성 평가",
        "bootstrap 95% 구간이 0을 포함했고, permutation test도 비유의였다.",
        "Gemini의 AP 우위는 안정적인 차이로 보기 어렵다.",
    ],
    [
        "6차년도 단일 wave",
        "전문가 어휘 기준 AP가 Gemini보다 높았다.",
        "pooled 분석의 소폭 우위가 단일 시점에서는 재현되지 않았다.",
    ],
    [
        "TF-IDF 기준",
        "TF-IDF n-gram AP는 Gemini보다 낮았다.",
        "Gemini는 단순 문자열 유사도보다 강했지만, 전문가 어휘 기준은 넘지 못했다.",
    ],
    [
        "사분면 분석",
        "both-high 조합의 양성률이 가장 높았다(.538/.636).",
        "두 기준이 수렴하는 조합을 우선 검토하는 전략이 가장 방어 가능하다.",
    ],
    [
        "판단 근거 점검",
        "엄격 지시문은 설명 양식을 일부 바꾸었으나 위험 신호를 제거하지 못했다.",
        "자동 점검은 타당성 검증이 아니라 오류 가능성 점검으로 해석해야 한다.",
    ],
]

for r, row_values in enumerate(rows):
    for c, value in enumerate(row_values):
        cell = table.cell(r, c)
        cell.text = value
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
        set_cell_margins(cell)
        set_cell_width(cell, [1.65, 2.65, 3.05][c])
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            for run in p.runs:
                run.font.name = "맑은 고딕"
                run.font.size = Pt(9.2 if r else 9.5)
                if r == 0:
                    run.bold = True

doc.save(OUT)
print(OUT)
