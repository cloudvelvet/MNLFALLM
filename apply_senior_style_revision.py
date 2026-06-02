from __future__ import annotations

from pathlib import Path

from docx import Document


SRC = Path(r"C:\chen_bauer_2024\docx_work\senior_source.docx")
OUT = Path(r"C:\chen_bauer_2024\docx_work\senior_style_revision_clean.docx")


def set_text(p, text: str) -> None:
    for run in p.runs:
        run.text = ""
    if p.runs:
        p.runs[0].text = text
    else:
        p.add_run(text)


doc = Document(SRC)

updates = {
    4: "본 연구는 다문화청소년패널(MAPS) 2기 자료의 청소년 및 보호자 문항을 대상으로, Gemini 2.5 Flash가 순서형 DIF 검토를 위한 문항-공변량 후보를 전문가 어휘 기준보다 효과적으로 우선순위화할 수 있는지 평가하였다. 분석 단위는 문항-공변량 조합이었다. Gemini 2.5 Flash에는 문항 문장, 구인 맥락, 응답 범주, 공변량 정의를 제공하되, 경험적 DIF 선별 결과, p값, 효과크기, 문항 통계량, 어휘 기준 점수는 제공하지 않았다. 비교 기준은 연구자가 사전에 구성한 이분형 전문가 어휘 기준, TF-IDF n-gram 의미 유사도 기준, 그리고 순서형 DIF 선별 기준이었다.",
    5: "평가는 평균 정밀도(AP)와 상위 후보 정밀도(P@5, P@10)를 함께 사용하였다. 1-5차년도 pooled 주 분석에서 Gemini의 AP는 기본 지시문 .382, 엄격 지시문 .405였고, 이분형 전문가 어휘 기준의 평균 AP는 각각 .375와 .384였다. 그러나 문항 단위 bootstrap 구간은 0을 포함했고, paired permutation test에서도 Gemini의 우위는 통계적으로 뚜렷하지 않았다. 또한 엄격 지시문에서는 AP가 소폭 증가했지만 P@5와 P@10은 낮아졌으며, 6차년도 단일 wave 민감도 분석에서는 전문가 어휘 기준이 Gemini보다 높은 AP를 보였다. 따라서 본 연구의 결과는 Gemini가 전문가 어휘 기준을 안정적으로 능가했다는 결론을 지지하지 않는다.",
    6: "이 결과는 Gemini 2.5 Flash를 DIF 판정 도구로 사용하는 것을 지지하지 않는다. 본 연구의 핵심 결론은 생성형 언어모형의 일반적 우수성이 아니라, 단일 상용 LLM을 DIF 후보 생성에 사용할 때 나타나는 조건과 한계이다. Gemini는 문항-공변량 조합에 대한 구조화된 설명을 생성할 수 있었지만, 그 점수와 상위 후보 목록은 지시문 조건과 경험적 선별 기준에 민감하였다. 따라서 LLM 산출물은 단독 판단 근거가 아니라, 전문가 어휘 기준, 경험적 선별, 전문가 내용 검토, 후속 심리측정 모형과 함께 사용되는 보조적 가설 자료로 다루어져야 한다.",
    11: "The findings do not support using Gemini 2.5 Flash as a DIF decision tool. Conditional on the screening criterion used here, Gemini produced structured and reviewable hypotheses but did not show a stable advantage over the expert lexical benchmark. The results instead identify boundary conditions for LLM-assisted DIF candidate generation: prompt sensitivity, divergence between global AP and top-k utility, dependence on covariate type, and the need to treat rationale audit as exploratory error-pattern evidence rather than validated qualitative interpretation.",
    23: "본 연구는 MAPS 청소년 및 보호자 응답 문항을 대상으로, 순서형 로지스틱 회귀에 기반한 DIF 선별 결과를 비교 기준으로 삼아 LLM의 우선순위화 성능을 평가한다. 이 선별 결과는 최종 DIF 판정이나 참값이 아니라, LLM 산출물과 전문가 어휘 기준이 경험적 선별 신호와 어느 정도 정렬되는지 평가하기 위한 provisional criterion이다. 구체적인 연구 질문은 다음과 같다.",
    25: "연구질문 2: Gemini가 산출한 DIF 가능성 점수는 이분형 전문가 어휘 기준보다 경험 DIF 후보를 더 효율적으로 우선순위화하는가?",
    27: "연구질문 4: 실제 차이(impact)와 문항 기능 차이(DIF)의 구분을 명시적으로 제약한 엄격한 지시문(strict prompt)은 기본 지시문(base prompt)에 비해 Gemini의 우선순위화 성능과 상위 후보 목록의 안정성을 개선하는가? 판단 근거 분석은 타당성 검증이 아니라, 지시문 조건에 따라 설명에서 나타나는 오류 유형이 어떻게 달라지는지 살펴보는 보조 분석으로 제한한다.",
    51: "2.4 순서형 DIF 선별 기준",
    52: "LLM과 전문가 어휘 기준의 우선순위화 성능은 순서형 DIF 선별 결과와 비교하였다. 이 선별 결과는 최종 DIF 판정이나 참값(gold standard)이 아니다. 본 연구에서 이 기준은 LLM 산출물과 어휘 기반 비교 기준이 경험적 선별 신호와 어느 정도 정렬되는지 평가하기 위한 provisional criterion으로 사용된다.",
    56: "공변량 계수는 uniform ordinal DIF의 예비 신호로 해석하였다. 즉, 같은 척도 점수 proxy 수준에서 공변량에 따라 특정 문항 응답 범주의 선택 경향이 달라지는지를 본 것이다. 본 연구는 nonuniform DIF를 모형화하지 않았고, proportional-odds 가정도 문항별로 별도 검정하지 않았다. p값에는 Benjamini-Hochberg 방식의 FDR 보정을 적용하였다(Benjamini & Hochberg, 1995). FDR < .05이면서 공변량 계수의 절댓값이 .20 이상인 조합을 경험 DIF 후보로 표시하였다. 이 기준은 통계적 유의성만으로 과도하게 많은 후보가 발생하는 것을 줄이기 위한 practical nonzero 기준이다.",
    62: "2.6 LLM 판단 근거의 보조 오류 유형 점검",
    64: "이 절차는 판단 근거의 심리측정학적 타당성을 검증하기 위한 것이 아니라, LLM 설명에서 나타나는 위험 신호를 기술하기 위한 보조 점검이다. 자동 코딩의 해석 가능성을 점검하기 위해 116개 문항-공변량-지시문 조합을 층화 표집하여 수동 검증용 파일을 마련했지만, 독립 코딩이 완료되지 않은 상태에서는 자동 코딩 결과를 판단 근거의 타당성 증거로 해석하지 않는다. 따라서 본 연구에서 자동 코딩은 LLM 설명의 오류 유형을 요약하는 탐색적 절차로만 사용하였다.",
    75: "표 5는 주 분석에서 LLM 점수와 이분형 전문가 어휘 기준의 우선순위화 성능을 비교한 결과이다. 기본 지시문 조건에서 분석 가능한 조합은 486개였고, 이 중 경험 DIF 후보는 128개였다. 엄격 지시문 조건에서는 482개 조합 중 127개가 경험 DIF 후보로 분류되었다. 양성 비율은 두 조건 모두 약 .263이었다.",
    84: "성별에서는 기본 지시문 LLM이 전문가 어휘 기준보다 약간 높았으나, 엄격 지시문에서는 그 차이가 사라졌다. 가구소득은 경험 후보가 세 개뿐이어서 AP가 매우 불안정하므로 강한 해석을 피해야 한다.",
    88: "†가구소득의 양성 후보는 3개에 불과하여 AP 추정치의 신뢰도가 극히 낮다. 해당 수치는 다른 공변량 결과와 직접 비교하기 어려우며, 해석 시 주의가 필요하다.",
    92: "LLM 점수의 보정 가능성도 별도로 확인하였다(표 8). 점수가 높은 구간에서 선별 양성률이 대체로 높아지는 경향은 있었지만, LLM 점수 자체를 경험적 확률로 해석하기에는 차이가 컸다. 예를 들어 80-100점 구간의 양성률은 기본 지시문에서 .407, 엄격 지시문에서 .479였다. 따라서 본 연구에서 LLM 점수는 보정된 확률이 아니라 후보 우선순위화를 위한 순위 신호로 해석한다.",
    93: "표 8. LLM 점수 구간별 선별 양성률",
    100: "3.6 보조 오류 유형 점검 결과",
    104: "표 10. LLM 판단 근거의 보조 오류 유형 점검 결과",
    109: "3.7 예시적 후보 사례",
    111: "표 12. 예시적 후보 사례",
    115: "경험 DIF 선별 기준을 달리했을 때도 결론의 방향은 크게 달라지지 않았다(표 13). 기준을 완화하거나 강화하면 양성 조합 수와 양성 비율은 변하지만, LLM과 전문가 어휘 기준의 상대적 성능은 조건에 따라 달라졌다. 특히 효과크기 기준을 강하게 적용할수록 양성 조합 수가 줄어들기 때문에 AP의 안정성도 함께 낮아진다. 따라서 본 연구의 결과는 특정 임계값 하나에 의존한 확정적 결론이라기보다, 여러 기준에서 나타나는 성능 패턴으로 읽어야 한다.",
    140: "3.12 LLM-전문가 어휘 사분면과 표적 random-intercept probe",
    141: "LLM 점수와 전문가 어휘 기준의 수렴 여부를 보기 위해 네 사분면을 구성하였다(표 18). 두 기준이 모두 높은 조합(both-high)은 경험적 양성 비율이 가장 높았다. 기본 지시문에서 이 영역의 cluster-robust 양성률은 .538이었고, 엄격 지시문에서는 .636이었다. 반대로 두 기준이 모두 낮은 영역(both-low)의 cluster-robust 양성률은 기본 지시문 .144, 엄격 지시문 .112로 낮았다. 이는 두 기준이 수렴하는 영역이 경험적으로 더 강한 후보군을 형성할 수 있음을 보여준다.",
    142: "반면 LLM 점수만 높은 영역(LLM-high/lexical-low)의 양성률은 중간 수준이었다. 이는 LLM 단독 고득점이 일부 경험적 신호를 포함하지만, 전문가 어휘 기준과 수렴할 때만큼 강한 후보군을 만들지는 못한다는 점을 시사한다. 전문가 어휘 기준만 높은 영역(lexical-high/LLM-low)에서도 양성 사례가 적지 않았기 때문에, LLM 점수만으로 어휘 기반 후보를 배제하는 것은 적절하지 않다.",
    144: "사분면별 표적 사례 8개에 대해 random-intercept ordinal probe를 추가로 수행하였다(표 19). 모든 사례에서 모형이 정상 수렴하였고, cluster-robust 선별 결과와 random-intercept probe의 계수 방향은 8개 사례 모두에서 일치하였다. 선별 label은 7개 사례에서 일치하였다. 불일치한 한 사례는 cluster-robust 기준에서는 양성이었으나, random-intercept probe에서는 FDR 보정 후 기준을 통과하지 못하였다.",
    146: "표 19. 사분면별 표적 사례의 random-intercept ordinal probe 요약",
    147: "주. RI = random-intercept ordinal probe. 각 사분면에서 표적 사례 1개를 선정하였으며, 응답자 수가 큰 경우 고정 seed로 최대 400명까지 표집하였다. 이 분석은 전체 검증이 아니라 표적 사례 검토이다.",
    152: "특히 6차년도 단일 wave 민감도 분석에서는 전문가 어휘 기준의 AP가 Gemini를 상회하였다. 이는 1-5차년도 pooled 자료에서 관찰된 소폭의 AP 우위가 단일 시점 자료에서 재현되지 않았음을 의미한다. 다만 6차년도 결과 역시 선별 기준에 의존하므로, 이를 Gemini의 실패로 단정하기보다 후보 우선순위화 결과가 자료 구성과 criterion 정의에 민감하다는 증거로 해석하는 것이 적절하다.",
    153: "따라서 Gemini 점수와 선별 기준 사이의 불일치는 Gemini의 실패로 단정할 수 없다. 기준 자체의 위양성률과 위음성률이 불명확하기 때문에, 기준 노이즈와 모델 한계를 구분하는 것은 현재 설계로는 불가능하다.",
    157: "넷째, 판단 근거 분석은 본문 핵심 타당성 검증이 아니라 보조적 오류 유형 점검으로 해석되어야 한다. 엄격한 DIF 구분 지시문은 Gemini가 '동일한 잠재특성 수준' 조건이나 응답 과정 관련 표현을 더 자주 사용하도록 만들었다. 그러나 이러한 형식적 개선이 상위 후보 정밀도 개선으로 이어지지는 않았다. 또한 성별·연령에 관한 일반적 설명이나 실제 잠재특성 차이(impact)를 문항 기능 차이(DIF)처럼 서술하는 위험 신호가 여전히 관찰되었다. 따라서 본 연구의 판단 근거 분석은 설명의 타당성을 입증하는 결과가 아니라, LLM 설명을 사용할 때 확인해야 할 오류 유형을 탐색적으로 요약한 결과로 보아야 한다.",
    160: "실무적 관점에서 본 연구의 결과는 Gemini를 단독적인 DIF 탐색 도구로 사용하는 것을 지지하지 않는다. 보다 방어 가능한 활용 방식은 다단계 절차이다. 먼저 전문가 어휘 기준으로 표면 단서가 분명한 조합을 확인하고, 그 다음 Gemini 산출물을 어휘 단서가 약하지만 의미적 관련성이 의심되는 조합을 설명적으로 정리하는 보조 자료로 사용한다. 마지막으로 high-risk 후보는 순서형 DIF 모형, IRT 기반 절차, MNLFA 등 별도의 심리측정 모형과 전문가 내용 검토를 통해 확인해야 한다. 특히 두 기준이 동시에 높은 조합(both-high)은 경험적 선별 양성률이 가장 높았으므로(.538/.636), 이를 최우선 전문가 검토 대상으로 설정하는 실무 절차를 권장한다.",
    161: "4.3 해석 범위와 후속 과제",
    162: "본 연구의 결론은 Gemini 2.5 Flash라는 단일 폐쇄형 모델, 두 가지 지시문 조건, 그리고 MAPS 2기 문항 pool에 한정된다. 이 범위는 작지 않지만, LLM 일반의 능력을 말하기에는 충분하지 않다. 또한 API 호출은 공개 모델 문자열(gemini-2.5-flash)을 사용하여 수행되었고, 제공자가 고정 snapshot 버전 또는 장기 보존 가능한 모델 해시를 제공하지 않는 조건에서 이루어졌다. 따라서 동일한 지시문과 입력을 사용하더라도 향후 API 업데이트 이후에는 완전히 동일한 산출물이 재현되지 않을 수 있다.",
    163: "문화·언어적 맥락도 별도의 검증을 요구한다. Gemini 2.5 Flash가 한국 다문화 청소년과 보호자의 사회문화적 맥락을 어느 정도 적절히 추론하는지는 본 연구만으로 확인할 수 없다. 이중문화 정체성, 이주 가정 경험, 한국어 사용 맥락은 문항의 표면 의미만으로 충분히 환원되기 어렵다. 주로 다국어·영어권 코퍼스를 포함해 훈련된 폐쇄형 모델이 이러한 맥락을 어떤 방식으로 일반화하는지는 후속 연구에서 따로 다루어야 한다.",
    164: "선별 기준의 한계는 결과 해석의 가장 중요한 제약이다. 본 연구의 proportional-odds 누적로짓 모형은 문항별 비례오즈 가정을 별도로 검정하지 않았고, pooled long format은 반복 측정의 종단적 의존성을 완전히 모형화하지 않는다. leave-one-item-out 척도 점수 proxy도 잠재특성의 직접 추정치가 아니므로, 척도의 차원성이나 문항 수에 따라 matching variable의 정밀도가 달라질 수 있다. 이 점 때문에 본 연구는 Gemini의 우위나 실패를 확정하지 않고, 기준에 조건화된 우선순위화 성능만 보고한다.",
    165: "마지막으로 판단 근거 점검은 설명의 타당성을 검증하는 절차가 아니다. 자동 코딩은 LLM 설명에서 나타나는 위험 신호의 분포를 요약하는 데 유용하지만, 독립적인 전문가 코딩이나 인지면담 자료를 대체하지 않는다. 따라서 rationale 결과는 LLM 설명의 품질을 확정하는 근거가 아니라, 후속 검토에서 주의해야 할 오류 유형을 정리한 보조 자료로 해석해야 한다.",
}

for idx, text in updates.items():
    if idx < len(doc.paragraphs):
        set_text(doc.paragraphs[idx], text)

replacements = {
    "잠정적 경험 DIF 후보": "경험 DIF 후보",
    "잠정적 경험 후보": "경험 후보",
    "잠정적 선별 양성률": "선별 양성률",
    "잠정적 경험 기준": "선별 기준",
    "잠정적 기준": "선별 기준",
    "잠정적 양성 후보": "양성 후보",
    "잠정적 경험 DIF 선별 기준": "경험 DIF 선별 기준",
    "잠정적 순서형 DIF 선별 결과": "순서형 DIF 선별 결과",
    "잠정적 DIF 선별 결과": "DIF 선별 결과",
    "failure-mode audit": "오류 유형 점검",
    "failure-mode 신호": "오류 유형 신호",
    "focal 후보 사례": "예시적 후보 사례",
    "Illustrative focal 후보 사례": "예시적 후보 사례",
    "focal 사례": "표적 사례",
}

for p in doc.paragraphs:
    text = p.text
    new = text
    for old, repl in replacements.items():
        new = new.replace(old, repl)
    if new != text:
        set_text(p, new)

for tbl in doc.tables:
    for row in tbl.rows:
        for cell in row.cells:
            text = cell.text
            new = text
            for old, repl in replacements.items():
                new = new.replace(old, repl)
            if new != text:
                cell.text = new

for p in list(doc.paragraphs):
    if p.text.strip() == ".":
        p._element.getparent().remove(p._element)

doc.save(OUT)
print(OUT)
