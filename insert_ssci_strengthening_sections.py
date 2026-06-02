from pathlib import Path
import csv
import re

from docx import Document
from docx.shared import Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


FONT = "맑은 고딕"
ROOT = Path(r"C:\chen_bauer_2024\MNLFALLM")
OUT = ROOT / "llm_dif_output"


def read_csv(path):
    with open(path, newline="", encoding="utf-8-sig") as f:
        return list(csv.DictReader(f))


def fmt(x):
    x = float(x)
    s = f"{x:.3f}"
    return s.replace("0.", ".").replace("-0.", "-.")


def set_run(run, size=10.5, bold=None):
    run.font.name = FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold


def style_paragraph(p, kind="body"):
    if kind == "h2":
        p.style = "Heading 2"
        size, bold = 13, True
    elif kind == "h3":
        p.style = "Heading 3"
        size, bold = 11, True
    elif kind == "caption":
        p.style = "Caption"
        size, bold = 9.5, False
    else:
        p.style = "Normal"
        size, bold = 10.5, False
    if not p.runs:
        p.add_run("")
    for r in p.runs:
        set_run(r, size, bold)


def insert_paragraph_before(target, text, kind="body"):
    new_p = OxmlElement("w:p")
    target._p.addprevious(new_p)
    p = target._parent.add_paragraph()
    p._p = new_p
    p.add_run(text)
    style_paragraph(p, kind)
    return p


def insert_table_before(doc, target, rows):
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    for i, row in enumerate(rows):
        for j, value in enumerate(row):
            cell = table.cell(i, j)
            cell.text = str(value)
            for p in cell.paragraphs:
                for r in p.runs:
                    set_run(r, 8.5, True if i == 0 else None)
    target._p.addprevious(table._tbl)
    return table


def delete_paragraph(p):
    elem = p._element
    elem.getparent().remove(elem)
    p._p = p._element = None


def find_para(doc, text):
    for p in doc.paragraphs:
        if p.text.strip() == text:
            return p
    raise ValueError(f"paragraph not found: {text}")


def remove_old_discussion(doc):
    start = None
    end = None
    for i, p in enumerate(doc.paragraphs):
        if p.text.strip() == "4. 논의":
            start = i
        if p.text.strip() == "References":
            end = i
            break
    if start is None or end is None or start >= end:
        return
    for p in list(doc.paragraphs[start:end]):
        delete_paragraph(p)


def main():
    src = max(Path("N:/").glob("개인/다음논문/*글자크기정리.docx"), key=lambda p: p.stat().st_mtime)
    dst = src.with_name(src.stem + "_SSCI보강.docx")
    doc = Document(src)

    # Add a method subsection before Results.
    results_p = find_para(doc, "3. 결과")
    method_sections = [
        ("2.7 불확실성 평가와 보조 경험 검증", "h3"),
        ("LLM과 비교 기준의 차이가 안정적인지 확인하기 위해 문항 단위 cluster bootstrap과 paired permutation test를 추가로 수행하였다. Bootstrap에서는 문항 ID를 복원추출 단위로 삼아 선택된 문항에 해당하는 모든 문항-공변량 조합을 함께 포함하였다. 각 반복에서 LLM AP, 이분형 키워드 AP, 두 AP의 차이를 계산하고 percentile 방식으로 95% 구간을 산출하였다. Permutation test에서는 같은 문항-공변량 조합 안에서 LLM 점수와 키워드 점수의 배정을 무작위로 교환하여 관찰된 AP 차이가 우연적 배정 아래에서 어느 정도 나타나는지 평가하였다.", "body"),
        ("또한 단순 어휘 기준과 LLM 사이에 또 다른 비교 기준을 두기 위해 API를 사용하지 않는 의미 유사도 기준을 구성하였다. 로컬 환경에 sentence-transformer 캐시가 없어, 본 연구에서는 문항 텍스트와 공변량 정의 사이의 TF-IDF n-gram cosine similarity를 산출하였다. 이 기준은 심층 의미모형이라기보다, 표면 단어 일치보다 넓은 문자열 기반 유사도를 포착하는 투명한 비전문가 기준으로 해석한다.", "body"),
        ("마지막으로 LLM 점수와 키워드 기준이 수렴하거나 불일치하는 조합의 경험적 양상을 검토하기 위해 사분면 분석을 수행하였다. LLM-high는 LLM 점수 70점 이상, LLM-low는 30점 이하로 정의하였고, keyword-high는 해당 공변량 키워드가 하나 이상 포함된 경우로 정의하였다. 이에 따라 LLM-high/keyword-low, keyword-high/LLM-low, both-high, both-low의 네 영역을 구성하였다. 각 영역에 대해 주 선별 기준과 cluster-robust 선별 기준의 양성 비율을 비교하였고, 각 영역의 대표 사례에 대해서는 random-intercept ordinal probe를 추가로 수행하였다. 이 probe는 전체 MNLFA가 아니라 focal triangulation이며, 응답자별 무선절편을 포함한 누적로짓 모형으로 계수 방향과 선별 label의 안정성을 점검하기 위한 보조 분석이다.", "body"),
    ]
    for text, kind in method_sections:
        insert_paragraph_before(results_p, text, kind)

    # Remove old discussion/limitations/conclusion before adding new result sections.
    remove_old_discussion(doc)
    ref_p = find_para(doc, "References")

    uncertainty = read_csv(OUT / "maps_llm_uncertainty_bootstrap_permutation.csv")
    uncertainty_rows = [["분석", "지시문", "LLM AP", "키워드 AP", "AP 차이", "Bootstrap 95% 구간", "Permutation p"]]
    prompt_label = {"original": "기본", "strict_dif": "엄격"}
    analysis_label = {"pooled_w1_w5": "1-5차 pooled", "wave6": "6차년도"}
    for r in uncertainty:
        uncertainty_rows.append([
            analysis_label[r["analysis"]],
            prompt_label[r["prompt_version"]],
            fmt(r["llm_ap"]),
            fmt(r["keyword_ap"]),
            fmt(r["delta_ap"]),
            f"[{fmt(r['delta_ci_low'])}, {fmt(r['delta_ci_high'])}]",
            fmt(r["perm_p_two_sided"]),
        ])

    embed_main = {r["prompt_version"]: r for r in read_csv(OUT / "maps_llm_embedding_eval_metrics.csv") if r["group"] == "overall"}
    embed_w6 = {r["prompt_version"]: r for r in read_csv(OUT / "maps_llm_embedding_eval_metrics_w6.csv") if r["group"] == "overall"}
    embed_rows = [["분석", "지시문", "LLM AP", "TF-IDF 유사도 AP", "LLM P@10", "TF-IDF P@10"]]
    for analysis, source in [("1-5차 pooled", embed_main), ("6차년도", embed_w6)]:
        for prompt in ["original", "strict_dif"]:
            r = source[prompt]
            embed_rows.append([
                analysis,
                prompt_label[prompt],
                fmt(r["llm_average_precision"]),
                fmt(r["embedding_average_precision"]),
                fmt(r["llm_precision_at_10"]),
                fmt(r["embedding_precision_at_10"]),
            ])

    quad = read_csv(OUT / "maps_llm_quadrant_validation_summary.csv")
    quad_rows = [["지시문", "사분면", "n", "주 선별 양성률", "Cluster-robust 양성률", "Cluster p 중앙값"]]
    for r in quad:
        quad_rows.append([
            prompt_label[r["prompt_version"]],
            r["quadrant"],
            r["n"],
            r["main_positive_rate"],
            r["cluster_positive_rate"],
            r["median_p_fdr_cluster"],
        ])

    ri = read_csv(OUT / "maps_dif_random_intercept_quadrant_probe_summary.csv")
    ri_rows = [["지시문", "사분면", "사례 수", "RI 수렴", "Cluster 양성", "RI 양성", "방향 일치율", "Label 일치율"]]
    for r in ri:
        ri_rows.append([
            prompt_label[r["prompt_version"]],
            r["quadrant"],
            r["n_cases"],
            r["ok_cases"],
            r["cluster_positive"],
            r["ri_positive"],
            fmt(r["sign_agree_rate"]),
            fmt(r["label_agree_rate"]),
        ])

    result_sections = [
        ("3.9 불확실성 평가", "h3"),
        ("문항 단위 cluster bootstrap과 paired permutation test를 추가한 결과, 주 분석에서 관찰된 LLM의 소폭 우위는 안정적인 차이로 해석되기 어려웠다(표 14). 1-5차년도 pooled 분석에서 기본 지시문의 AP 차이는 .007이었고 bootstrap 95% 구간은 [-.070, .071]이었다. 엄격 지시문에서도 AP 차이는 .020이었으나 bootstrap 95% 구간은 [-.063, .077]로 0을 포함하였다. Permutation test에서도 두 조건 모두 차이가 통계적으로 두드러지지 않았다.", "body"),
        ("6차년도 단일 wave 분석에서는 LLM의 AP가 키워드 기준보다 낮았지만, 이 차이 역시 bootstrap 구간에서 0을 포함하였다. 따라서 본 연구의 결과는 LLM의 안정적 예측 우위를 보여준다기보다, LLM과 투명한 키워드 기준 사이의 차이가 자료 구성과 선별 기준에 따라 흔들릴 수 있음을 보여준다.", "body"),
        ("표 14. LLM과 키워드 기준 AP 차이에 대한 bootstrap 및 permutation 결과", "caption"),
    ]
    for text, kind in result_sections:
        insert_paragraph_before(ref_p, text, kind)
    insert_table_before(doc, ref_p, uncertainty_rows)
    insert_paragraph_before(ref_p, "주. Bootstrap 단위는 문항 ID이며, AP 차이는 LLM AP에서 이분형 키워드 AP를 뺀 값이다. 키워드 AP는 동점 무작위 해소 평균을 사용하였다.", "caption")

    for text, kind in [
        ("3.10 의미 유사도 기준과의 비교", "h3"),
        ("추가 비교 기준으로 사용한 TF-IDF n-gram cosine similarity는 전체 AP에서 LLM보다 낮았다(표 15). 1-5차년도 pooled 분석에서 의미 유사도 기준의 AP는 기본 지시문 .291, 엄격 지시문 .292였고, LLM AP는 각각 .382와 .405였다. 6차년도 단일 wave에서도 의미 유사도 기준은 LLM보다 낮았다. 이는 LLM이 단순 문자열 기반 의미 유사도보다는 더 많은 후보 정보를 포착했음을 시사한다.", "body"),
        ("다만 이 결과를 LLM의 일반적 우수성으로 해석해서는 안 된다. 의미 유사도 기준은 투명한 비전문가 비교 기준이며, 이분형 키워드 기준은 일부 공변량에서 더 강하게 작동하였다. 따라서 본 연구의 비교 구조는 LLM이 약한 기준을 이겼다는 단순 결론이 아니라, LLM의 위치가 키워드 기준과 일반 유사도 기준 사이에서 어떻게 달라지는지를 보여준다.", "body"),
        ("표 15. TF-IDF 의미 유사도 기준과 LLM의 우선순위화 성능", "caption"),
    ]:
        insert_paragraph_before(ref_p, text, kind)
    insert_table_before(doc, ref_p, embed_rows)

    for text, kind in [
        ("3.11 LLM-키워드 사분면과 focal random-intercept probe", "h3"),
        ("LLM 점수와 키워드 기준의 수렴 여부를 보기 위해 네 사분면을 구성하였다(표 16). 두 기준이 모두 높은 both-high 영역은 경험적 양성 비율이 가장 높았다. 기본 지시문에서 both-high의 cluster-robust 양성률은 .538이었고, 엄격 지시문에서는 .636이었다. 반대로 both-low 영역의 cluster-robust 양성률은 기본 지시문 .144, 엄격 지시문 .112로 낮았다. 이는 LLM과 키워드 기준이 수렴하는 영역이 경험적으로 더 강한 후보군을 형성할 수 있음을 보여준다.", "body"),
        ("반면 LLM-high/keyword-low 영역의 양성률은 중간 수준이었다. 이는 LLM 단독 고득점이 일부 경험적 신호를 포함하지만, 키워드 기준과 수렴할 때만큼 강한 후보군을 만들지는 못한다는 점을 시사한다. keyword-high/LLM-low 영역에서도 양성 사례가 적지 않았기 때문에, LLM 점수만으로 키워드 기반 후보를 배제하는 것은 적절하지 않다.", "body"),
        ("표 16. LLM 점수와 키워드 기준의 사분면별 경험적 선별 양상", "caption"),
    ]:
        insert_paragraph_before(ref_p, text, kind)
    insert_table_before(doc, ref_p, quad_rows)

    for text, kind in [
        ("사분면별 대표 사례 8개에 대해 random-intercept ordinal probe를 추가로 수행하였다(표 17). 모든 사례에서 모형이 정상 수렴하였고, cluster-robust 선별 결과와 random-intercept probe의 계수 방향은 8개 사례 모두에서 일치하였다. 선별 label은 7개 사례에서 일치하였다. 불일치한 한 사례는 cluster-robust 기준에서는 양성이었으나, random-intercept probe에서는 FDR 보정 후 기준을 통과하지 못하였다.", "body"),
        ("이 focal probe는 최종 검증이 아니라 표적 경험 검토이다. 그럼에도 cluster-robust 결과와 random-intercept probe의 방향성이 대체로 일치했다는 점은, 본 연구의 선별 label이 전적으로 표준오차 처리 방식에만 의존한 산물이 아님을 보여준다. 동시에 일부 후보의 경험적 지지는 모형 구조에 따라 약화될 수 있으므로, LLM 산출물은 여전히 검증할 후보로 다루어야 한다.", "body"),
        ("표 17. 사분면별 대표 사례의 random-intercept ordinal probe 요약", "caption"),
    ]:
        insert_paragraph_before(ref_p, text, kind)
    insert_table_before(doc, ref_p, ri_rows)
    insert_paragraph_before(ref_p, "주. RI = random-intercept ordinal probe. 각 사분면에서 대표 사례 1개를 선정하였으며, 응답자 수가 큰 경우 고정 seed로 최대 400명까지 표집하였다. 이 분석은 전체 검증이 아니라 focal triangulation이다.", "caption")

    discussion_texts = [
        ("4. 논의", "h2"),
        ("본 연구는 MAPS 문항-공변량 조합을 대상으로, 경험적 DIF 선별 결과를 보지 않은 LLM이 후속 검토를 위한 후보 가설을 생성하고 우선순위화할 수 있는지 평가하였다. 새롭게 추가한 불확실성 평가, 의미 유사도 기준, 사분면 분석, focal random-intercept probe를 종합하면, 본 연구의 핵심 결과는 LLM의 강한 예측 우위가 아니라 LLM 기반 후보 생성의 경계 조건을 경험적으로 드러낸 데 있다.", "body"),
        ("첫째, LLM은 구조화된 후보 가설과 판단 근거를 생성할 수 있었다. 이는 문항 문장, 구인 맥락, 응답 범주, 공변량 정의를 함께 읽고 검토 가능한 설명을 산출할 수 있음을 보여준다. 그러나 이 산출물은 DIF 판정이 아니라 검증을 기다리는 후보 설명이다.", "body"),
        ("둘째, LLM의 단독 우위는 불안정하였다. 1-5차년도 pooled 분석의 전체 AP에서는 LLM이 이분형 키워드 기준을 소폭 상회했지만, bootstrap 구간은 0을 포함했고 permutation test에서도 뚜렷한 차이가 확인되지 않았다. 6차년도 단일 wave 분석에서는 오히려 키워드 기준이 높았다. 따라서 LLM 점수는 독립적인 판정 신호라기보다 조건부 우선순위화 신호로 이해하는 것이 적절하다.", "body"),
        ("셋째, LLM은 TF-IDF 기반 의미 유사도 기준보다는 높은 성능을 보였다. 이는 LLM이 단순 문자열 유사도 이상의 정보를 일부 포착했음을 시사한다. 그러나 한국어 능력이나 연령처럼 표면 단서가 강하게 작동하는 공변량에서는 키워드 기준이 더 안정적으로 기능하였다. 따라서 LLM의 가치는 단순 baseline을 대체하는 데 있지 않고, 여러 기준과 함께 후보 공간을 분해하는 데 있다.", "body"),
        ("넷째, LLM과 키워드 기준이 수렴하는 both-high 영역은 경험적으로 가장 강한 후보군을 형성하였다. 반대로 both-low 영역은 양성률이 낮았다. 이 결과는 LLM과 키워드 기준을 경쟁 관계로만 볼 필요가 없음을 보여준다. 두 기준이 서로 다른 단서를 제공하고, 그 수렴과 불일치를 이용해 후속 검토의 우선순위를 더 정교하게 설정할 수 있다.", "body"),
        ("마지막으로 focal random-intercept probe는 cluster-robust 선별 결과와 대체로 같은 방향의 경험적 신호를 보였다. 이는 본 연구의 경험적 후보 label이 완전히 취약한 산물이 아님을 보여주지만, 동시에 일부 사례에서는 모형 구조가 달라지면 label이 약화될 수 있음을 확인하였다. 따라서 본 연구의 결론은 LLM이 DIF를 탐지한다는 주장이 아니라, LLM 산출물을 투명한 baseline, 선별 label의 민감도, 표적 경험 검증과 함께 배치해야 한다는 주장이다.", "body"),
    ]
    for text, kind in discussion_texts:
        insert_paragraph_before(ref_p, text, kind)

    limitations = [
        ("5. 한계", "h2"),
        ("본 연구의 결과는 몇 가지 제한 안에서 해석되어야 한다. 첫째, 잠정적 순서형 DIF 선별 양성 조합은 DIF의 참값이 아니라 경험적 검토 후보이다. 본 연구는 cluster-robust 분석, 6차년도 단일 wave 분석, bootstrap, permutation, focal random-intercept probe를 통해 민감도를 점검하였지만, 전체 MNLFA나 완전한 종단 측정동일성 모형을 추정한 것은 아니다.", "body"),
        ("둘째, leave-one-item-out 척도 점수 proxy는 잠재특성의 직접 추정치가 아니다. 척도 구조가 다차원적이거나 문항 수가 적은 경우 matching variable의 한계가 커질 수 있으며, 일부 문항이 DIF를 포함할 경우 anchor contamination 가능성도 남아 있다.", "body"),
        ("셋째, 본 연구는 단일 LLM인 Gemini 2.5 Flash와 두 지시문 조건에 기반한다. 폐쇄형 LLM은 모델 버전과 API 설정이 시간에 따라 변할 수 있으므로, 본 연구의 결론은 특정 시점의 특정 모델과 지시문 조건에 대한 탐색적 비교 결과로 해석해야 한다.", "body"),
        ("넷째, 판단 근거 코딩은 자동 규칙 기반으로 수행되었다. 자동 코딩은 LLM 설명의 양식과 위험 신호를 요약하는 audit 절차이지, 판단 근거의 심리측정학적 타당성을 입증하는 절차가 아니다. 독립 전문가 코딩이 추가된다면 LLM 설명의 실제 검토 가능성을 더 직접적으로 평가할 수 있다.", "body"),
        ("다섯째, 키워드 기준과 TF-IDF 의미 유사도 기준은 모두 투명한 비교 기준이지만 중립적 기준은 아니다. 키워드 목록은 연구자의 판단에 의해 구성되었고, TF-IDF 유사도는 문맥적 추론을 충분히 반영하지 못한다. 따라서 본 연구의 baseline은 LLM을 평가하기 위한 비교 장치이지 최종적 대안 모형이 아니다.", "body"),
    ]
    for text, kind in limitations:
        insert_paragraph_before(ref_p, text, kind)

    conclusion = [
        ("6. 결론", "h2"),
        ("결론적으로, 본 연구는 LLM을 DIF 판정 도구로 지지하지 않는다. LLM은 일부 조건에서 키워드 기준이나 TF-IDF 유사도 기준보다 높은 우선순위화 성능을 보였지만, 그 우위는 bootstrap, permutation, 6차년도 민감도 분석에서 안정적으로 확인되지 않았다. 따라서 LLM의 역할은 자동 탐지가 아니라 후보 가설 생성과 검토 순서 제안으로 제한되어야 한다.", "body"),
        ("본 연구의 기여는 LLM이 DIF를 잘 찾아낸다는 주장에 있지 않다. 오히려 LLM 산출물을 후보 가설로 제한하고, 키워드 기준, 의미 유사도 기준, prompt sensitivity, 경험적 선별 label의 민감도, 사분면 분석, focal random-intercept probe를 함께 놓고 평가하는 절차를 제안했다는 데 있다. 이 절차는 대규모 문항-공변량 공간에서 검토할 후보를 좁히는 데 활용될 수 있으며, 동시에 LLM이 어디에서 심리측정적 추론과 어긋나는지를 드러내는 안전장치가 될 수 있다.", "body"),
    ]
    for text, kind in conclusion:
        insert_paragraph_before(ref_p, text, kind)

    # Re-apply table font sizes to all tables, including inserted ones.
    for table in doc.tables:
        for row_i, row in enumerate(table.rows):
            for cell in row.cells:
                for p in cell.paragraphs:
                    for r in p.runs:
                        set_run(r, 8.5, True if row_i == 0 else None)

    doc.save(dst)
    print(dst)


if __name__ == "__main__":
    main()
