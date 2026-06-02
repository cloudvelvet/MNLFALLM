from pathlib import Path
import shutil

from docx import Document
from docx.shared import RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


WORK = Path(r"C:\chen_bauer_2024\docx_work")
SRC = WORK / "submission_trim_publication_ready.docx"
LOCAL_OUT = WORK / "submission_trim_publication_ready_black.docx"
FINAL_OUT = Path(r"N:\개인\다음논문\LLM_DIF_full_manuscript_draft_psychometric_style_Gemini범위수정_투고압축본_최종.docx")


def force_run_black(run):
    run.font.color.rgb = RGBColor(0, 0, 0)
    r_pr = run._r.get_or_add_rPr()
    color = r_pr.find(qn("w:color"))
    if color is None:
        color = OxmlElement("w:color")
        r_pr.append(color)
    color.set(qn("w:val"), "000000")
    if qn("w:themeColor") in color.attrib:
        del color.attrib[qn("w:themeColor")]


doc = Document(SRC)

for style in doc.styles:
    try:
        style.font.color.rgb = RGBColor(0, 0, 0)
    except Exception:
        pass

for paragraph in doc.paragraphs:
    for run in paragraph.runs:
        force_run_black(run)

for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    force_run_black(run)

doc.save(LOCAL_OUT)
shutil.copy2(LOCAL_OUT, FINAL_OUT)
print(LOCAL_OUT)
print(FINAL_OUT)
