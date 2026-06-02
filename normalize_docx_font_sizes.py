from pathlib import Path
import re

from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn


FONT = "맑은 고딕"


def set_run_font(run, size_pt, bold=None):
    run.font.name = FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    run.font.size = Pt(size_pt)
    if bold is not None:
        run.font.bold = bold


def set_para_font(paragraph, size_pt, bold=None):
    if not paragraph.runs:
        paragraph.add_run("")
    for run in paragraph.runs:
        set_run_font(run, size_pt, bold)


def set_style_font(doc, style_name, size_pt, bold=None):
    if style_name not in doc.styles:
        return
    style = doc.styles[style_name]
    style.font.name = FONT
    style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    style.font.size = Pt(size_pt)
    if bold is not None:
        style.font.bold = bold


def main():
    src = max(
        Path("N:/").glob("개인/다음논문/LLM_DIF_full_manuscript_draft_psychometric_style_*방법결과반영.docx"),
        key=lambda p: p.stat().st_mtime,
    )
    dst = src.with_name(src.stem + "_글자크기정리.docx")
    doc = Document(src)

    set_style_font(doc, "Normal", 10.5, False)
    set_style_font(doc, "Title", 16, True)
    set_style_font(doc, "Subtitle", 11, False)
    set_style_font(doc, "Heading 1", 13, True)
    set_style_font(doc, "Heading 2", 13, True)
    set_style_font(doc, "Heading 3", 11, True)
    set_style_font(doc, "Caption", 9.5, False)

    for p in doc.paragraphs:
        text = p.text.strip()
        if not text:
            continue
        if p.style.name == "Title":
            set_para_font(p, 16, True)
        elif p.style.name == "Subtitle":
            set_para_font(p, 11, False)
        elif re.match(r"^\d+\.\s", text) or text in {"국문초록", "English Abstract", "References"}:
            p.style = doc.styles["Heading 2"]
            set_para_font(p, 13, True)
        elif re.match(r"^\d+\.\d+\s", text):
            p.style = doc.styles["Heading 3"]
            set_para_font(p, 11, True)
        elif text.startswith("표 ") or text.startswith("주."):
            p.style = doc.styles["Caption"]
            set_para_font(p, 9.5, False)
        else:
            if p.style.name == "Caption":
                p.style = doc.styles["Normal"]
            set_para_font(p, 10.5, False)

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    for run in p.runs:
                        set_run_font(run, 8.5, None)

    # Make header rows a little clearer without changing table size.
    for table in doc.tables:
        if table.rows:
            for cell in table.rows[0].cells:
                for p in cell.paragraphs:
                    for run in p.runs:
                        set_run_font(run, 8.5, True)

    doc.save(dst)
    print(dst)


if __name__ == "__main__":
    main()
