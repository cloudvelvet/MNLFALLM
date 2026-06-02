from pathlib import Path
import shutil

from docx import Document
from docx.shared import RGBColor


WORK = Path(r"C:\chen_bauer_2024\docx_work")
SRC = WORK / "submission_trim_final2.docx"
LOCAL_OUT = WORK / "submission_trim_publication_ready.docx"
FINAL_OUT = Path(r"N:\개인\다음논문\LLM_DIF_full_manuscript_draft_psychometric_style_Gemini범위수정_투고압축본_최종.docx")


def blacken_body_runs(paragraph):
    text = paragraph.text.strip()
    style = paragraph.style.name if paragraph.style is not None else ""
    keep_color = style.startswith("Heading") or text.startswith("표 ")
    if keep_color:
        return
    for run in paragraph.runs:
        run.font.color.rgb = RGBColor(0, 0, 0)


doc = Document(SRC)

for paragraph in doc.paragraphs:
    blacken_body_runs(paragraph)

for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.color.rgb = RGBColor(0, 0, 0)

doc.save(LOCAL_OUT)
shutil.copy2(LOCAL_OUT, FINAL_OUT)
print(LOCAL_OUT)
print(FINAL_OUT)
