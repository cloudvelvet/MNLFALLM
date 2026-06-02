from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK


SRC = Path(r"C:\chen_bauer_2024\docx_work\current_for_required.docx")
OUT = Path(r"C:\chen_bauer_2024\docx_work\required_publication_fixes.docx")


def set_text(paragraph, text: str) -> None:
    for run in paragraph.runs:
        run.text = ""
    if paragraph.runs:
        paragraph.runs[0].text = text
    else:
        paragraph.add_run(text)


def add_after(paragraph, text: str, style=None):
    from docx.oxml import OxmlElement
    from docx.text.paragraph import Paragraph

    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    p = Paragraph(new_p, paragraph._parent)
    if style is not None:
        p.style = style
    p.add_run(text)
    return p


def insert_after_text(doc, startswith: str, text: str) -> bool:
    for p in doc.paragraphs:
        if p.text.strip().startswith(startswith):
            add_after(p, text, style=p.style)
            return True
    return False


doc = Document(SRC)

# Required-2: repair residual encoding and target-probe wording.
for p in doc.paragraphs:
    t = p.text.strip()
    if t.startswith("3.12"):
        set_text(p, "3.12 LLM-전문가 어휘 사분면과 focal random-intercept probe")
    elif t.startswith("이 focal probe는 최종 검증이 아니라"):
        set_text(
            p,
            "이 probe는 전체 MNLFA가 아니라 표적 경험 검토이다. 그럼에도 cluster-robust 결과와 random-intercept probe의 방향성이 대체로 일치했다는 점은, 본 연구의 선별 label이 전적으로 표준오차 처리 방식에만 의존한 산물이 아님을 보여준다. 동시에 일부 후보의 경험적 지지는 모형 구조에 따라 약화될 수 있으므로, LLM 산출물은 여전히 검증할 후보로 다루어야 한다.",
        )
    elif t.startswith("주. RI = random-intercept ordinal probe"):
        set_text(
            p,
            "주. RI = random-intercept ordinal probe. 각 사분면에서 focal 사례 1개를 선정하였으며, 응답자 수가 큰 경우 고정 seed로 최대 400명까지 표집하였다. 이 분석은 전체 검증이 아니라 표적 사례 검토이다.",
        )

# Required-5: TF-IDF specification.
for p in doc.paragraphs:
    if p.text.strip().startswith("또한 단순 어휘 기준과 LLM 사이에 또 다른 비교 기준"):
        set_text(
            p,
            "또한 단순 어휘 기준과 LLM 사이에 또 다른 비교 기준을 두기 위해 API를 사용하지 않는 의미 유사도 기준을 구성하였다. TF-IDF n-gram cosine similarity는 정확한 어휘 일치보다 넓은 표면 단어 분포를 포착하되, 사전 학습된 언어모형을 필요로 하지 않는 투명한 비전문가 기준으로 포함하였다. 본 분석은 문항/구인 맥락과 공변량 정의 텍스트를 대상으로 단어 1-2그램과 문자 3-5그램 TF-IDF 벡터를 구성하고 cosine similarity를 산출하였다. 최소 문서 빈도는 1로 두었으며, 별도의 형태소 분석이나 동의어 확장 없이 공백 정리와 기본 문자열 정규화만 적용하였다.",
        )
        break

# Required-1 and recommendation-1: add criterion-noise caveat and AP/P@k mechanism.
insert_after_text(
    doc,
    "특히 6차년도 단일 wave 민감도 분석에서는",
    "따라서 Gemini 점수와 잠정적 기준 사이의 불일치는 Gemini의 실패로 단정할 수 없다. 잠정적 기준 자체의 위양성률과 위음성률이 불명확하기 때문에, 기준 노이즈와 모델 한계를 구분하는 것은 현재 설계로는 불가능하다.",
)
for p in doc.paragraphs:
    if p.text.strip().startswith("둘째, 지시문 조건에 따른 상위 후보군"):
        set_text(
            p,
            "둘째, 지시문 조건에 따른 상위 후보군(top-k)의 민감도는 Gemini 기반 DIF 후보 생성의 중요한 경계조건으로 확인되었다. 기본 지시문과 엄격 지시문 조건 간 Spearman 순위상관은 .571이었고, top-5와 top-10의 중복 비율은 모두 0이었다. 엄격 지시문에서 전체 AP가 소폭 개선되었음에도 P@5와 P@10이 낮아진 결과는, 지시문 강화가 전체 순위의 일부 신호를 조정할 수는 있지만 실제 검토 대상이 되는 최상위 목록을 안정화하지는 못했음을 보여준다. 이 분기는 점수 분포 변화와 관련된다. 기본 지시문 조건에서 LLM-high(70점 이상) 영역에 속한 조합은 138개였으나, 엄격 지시문 조건에서는 198개로 늘어났다. 즉 엄격 지시문은 고점수 후보군의 폭을 넓히면서 전체 AP를 일부 높였지만, 최상위 목록에서는 경험적 양성 후보의 밀도를 희석했을 가능성이 있다. 따라서 prompt engineering만으로 신뢰할 만한 후보 목록을 얻는다는 주장은 본 자료에서 지지되지 않는다.",
        )
        break

# Recommendation-2: both-high practical recommendation.
for p in doc.paragraphs:
    if p.text.strip().startswith("실무적 관점에서 본 연구의 결과는 Gemini를"):
        set_text(
            p,
            p.text
            + " 특히 LLM 점수와 전문가 어휘 기준이 동시에 높은 both-high 조합은 경험적 선별 양성률이 가장 높았으므로(.538/.636), 이를 최우선 전문가 검토 대상으로 설정하는 실무 절차를 권장한다.",
        )
        break

# Recommendation-3 and 4: cultural-linguistic and API reproducibility limitations.
insert_after_text(
    doc,
    "본 연구는 상용 폐쇄형 LLM(Gemini 2.5 Flash)을 사용하였기에",
    "또한 Gemini 2.5 Flash가 한국 다문화 청소년과 보호자의 사회문화적 맥락을 어느 정도 적절히 추론하는지는 별도의 검증이 필요하다. 이중문화 정체성, 이주 가정 경험, 한국어 사용 맥락은 문항의 표면 의미만으로 충분히 환원되기 어렵다. 주로 다국어·영어권 코퍼스를 포함해 훈련된 폐쇄형 모델이 이러한 맥락을 어떤 방식으로 일반화하는지는 본 연구 설계만으로 확인할 수 없다.",
)
insert_after_text(
    doc,
    "또한 Gemini 2.5 Flash가 한국 다문화 청소년과 보호자의 사회문화적 맥락",
    "API 호출은 공개 모델 문자열(gemini-2.5-flash)을 사용하여 수행하였고, 제공자가 고정 snapshot 버전 또는 장기 보존 가능한 모델 해시를 제공하지 않는 조건에서 이루어졌다. 따라서 동일한 지시문과 입력을 사용하더라도 향후 API 업데이트 이후에는 완전히 동일한 산출물이 재현되지 않을 수 있다.",
)

# Required-3: table 6 household-income dagger + note.
tbl = doc.tables[5]
for row_idx in [9, 10]:
    cell = tbl.cell(row_idx, 0)
    if "†" not in cell.text:
        cell.text = cell.text + "†"

# Add note after Table 6 caption or right after paragraph before next section.
inserted_note = False
for i, p in enumerate(doc.paragraphs):
    if p.text.strip().startswith("주. 공변량별 AP"):
        add_after(
            p,
            "†가구소득의 잠정적 양성 후보는 3개에 불과하여 AP 추정치의 신뢰도가 극히 낮다. 해당 수치는 다른 공변량 결과와 직접 비교하기 어려우며, 해석 시 주의가 필요하다.",
            style=p.style,
        )
        inserted_note = True
        break
if not inserted_note:
    for p in doc.paragraphs:
        if p.text.strip().startswith("3.4 응답자 유형별"):
            # Insert before the next section heading.
            from docx.oxml import OxmlElement
            from docx.text.paragraph import Paragraph

            new_p = OxmlElement("w:p")
            p._p.addprevious(new_p)
            np = Paragraph(new_p, p._parent)
            np.style = doc.styles["Normal"]
            np.add_run("†가구소득의 잠정적 양성 후보는 3개에 불과하여 AP 추정치의 신뢰도가 극히 낮다. 해당 수치는 다른 공변량 결과와 직접 비교하기 어려우며, 해석 시 주의가 필요하다.")
            break

# Recommendation-4: table 3 model version information.
t3 = doc.tables[2]
row = t3.add_row()
row.cells[0].text = "API 버전 고정 여부"
row.cells[1].text = "공개 모델 문자열 gemini-2.5-flash를 사용하였다. 호출 당시 별도의 고정 snapshot ID 또는 모델 해시는 제공받지 못했으므로, 향후 제공자 측 모델 업데이트에 따른 재현성 제한이 있다."

# Required-4: add literature context sentence and references.
for p in doc.paragraphs:
    if p.text.strip().startswith("최근 Transformer 기반 언어모형을 심리측정 문항 설계"):
        add_after(
            p,
            "이 문제는 넓게 보면 문항에 대한 사전 판단이 실제 경험적 지표와 얼마나 정렬되는가라는 오래된 측정 문제와도 연결된다. 예컨대 Sandoval과 Miille(1980)는 소수집단 문항 난이도 판단에서 전문가 판단의 정확성 문제를 다루었고, 최근에는 Gilardi 등(2023)이 텍스트 주석 과업에서 LLM 기반 판단의 신뢰도를 논의하였다. 이러한 연구들은 문항 또는 텍스트에 대한 사전 판단이 유용할 수 있지만, 그 판단을 경험적 기준과 분리하여 검증해야 함을 시사한다.",
            style=p.style,
        )
        break

existing = "\n".join(p.text for p in doc.paragraphs)
refs = [
    "Gilardi, F., Alizadeh, M., & Kubli, M. (2023). ChatGPT outperforms crowd workers for text-annotation tasks. *Proceedings of the National Academy of Sciences, 120*(30), e2305016120.",
    "Sandoval, J., & Miille, M. P. W. (1980). Accuracy of judgments of WISC-R item difficulty for minority groups. *Journal of Consulting and Clinical Psychology, 48*(2), 249-253.",
]
insert_anchor = None
for p in doc.paragraphs:
    if p.text.strip().startswith("신동훈"):
        insert_anchor = p
        break
if insert_anchor is None:
    insert_anchor = doc.paragraphs[-1]
for ref in reversed(refs):
    first_author = ref.split(",")[0]
    if first_author not in existing:
        from docx.oxml import OxmlElement
        from docx.text.paragraph import Paragraph

        new_p = OxmlElement("w:p")
        insert_anchor._p.addprevious(new_p)
        np = Paragraph(new_p, insert_anchor._parent)
        np.style = insert_anchor.style
        np.add_run(ref)

# Some remaining exact repairs.
for p in doc.paragraphs:
    if "???" in p.text:
        set_text(p, p.text.replace("???", "focal"))

doc.save(OUT)
print(OUT)
