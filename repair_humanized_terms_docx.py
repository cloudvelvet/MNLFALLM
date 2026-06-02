from pathlib import Path
import shutil

from docx import Document


WORK = Path(r"C:\chen_bauer_2024\docx_work")
SRC = WORK / "submission_trim_humanized_terms.docx"
LOCAL_OUT = WORK / "submission_trim_humanized_terms_repaired.docx"
FINAL_OUT = Path(r"N:\개인\다음논문\LLM_DIF_full_manuscript_draft_psychometric_style_Gemini범위수정_투고압축본_최종.docx")


doc = Document(SRC)

english_replacements = {
    "This study examined whether Gemini 2.5 Flash can prioritize item-covariate combinations": (
        "This study examined whether Gemini 2.5 Flash can prioritize item-covariate combinations "
        "for ordinal differential item functioning (DIF) review better than an expert lexical benchmark. "
        "The analysis used youth and caregiver items from the second Multicultural Adolescents Panel Study (MAPS). "
        "Gemini 2.5 Flash received item wording, construct context, response categories, and covariate definitions, "
        "but no empirical DIF screening results, p values, effect sizes, item statistics, or lexical benchmark scores. "
        "Its rankings were compared with a pre-specified binary expert lexical benchmark, a TF-IDF n-gram similarity "
        "benchmark, and an ordinal DIF screening criterion."
    ),
    "Average precision (AP) and top-candidate precision": (
        "Average precision (AP) and top-candidate precision (P@5 and P@10) were treated as complementary prioritization "
        "indices. In the pooled Wave 1-5 analysis, Gemini AP was .382 under the original prompt and .405 under the "
        "strict prompt, whereas the binary expert lexical benchmark yielded mean AP values of .375 and .384, respectively. "
        "However, item-cluster bootstrap intervals included zero, and paired permutation tests did not show a statistically "
        "clear advantage for Gemini. Under the strict prompt, AP improved slightly but P@5 and P@10 declined; in the Wave 6 "
        "single-wave sensitivity analysis, the lexical benchmark outperformed Gemini in AP."
    ),
    "The findings do not support using Gemini 2.5 Flash": (
        "The findings do not support using Gemini 2.5 Flash as a DIF decision tool. Conditional on the screening criterion "
        "used here, Gemini produced structured and reviewable hypotheses but did not show a stable advantage over the expert "
        "lexical benchmark. The results instead identify boundary conditions for LLM-assisted DIF candidate generation: "
        "prompt sensitivity, divergence between global AP and top-k utility, dependence on covariate type, and the need to "
        "treat rationale audit as exploratory error-pattern evidence rather than validated qualitative interpretation."
    ),
    "Keywords: differential item functioning": (
        "Keywords: differential item functioning, large language models, multicultural adolescents, lexical benchmark, "
        "prompt sensitivity, psychometric benchmarking"
    ),
}

for paragraph in doc.paragraphs:
    stripped = paragraph.text.strip()
    for prefix, new_text in english_replacements.items():
        if stripped.startswith(prefix):
            paragraph.text = new_text

for paragraph in doc.paragraphs:
    text = paragraph.text
    if "LLM-high/어휘-low, 어휘-high/LLM-low, both-high, both-low" in text:
        paragraph.text = text.replace(
            "LLM-high/어휘-low, 어휘-high/LLM-low, both-high, both-low",
            "LLM-high/어휘-low, 어휘-high/LLM-low, 두 기준 모두 높음(both-high), 두 기준 모두 낮음(both-low)",
        )
    if "두 기준이 동시에 높은 both-high 조합" in text:
        paragraph.text = text.replace(
            "두 기준이 동시에 높은 both-high 조합",
            "두 기준이 동시에 높은 조합(both-high)",
        )

for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                if "both-high 조합의 양성률" in paragraph.text:
                    paragraph.text = paragraph.text.replace(
                        "both-high 조합의 양성률",
                        "두 기준 모두 높음(both-high) 조합의 양성률",
                    )

doc.save(LOCAL_OUT)
shutil.copy2(LOCAL_OUT, FINAL_OUT)
print(LOCAL_OUT)
print(FINAL_OUT)
