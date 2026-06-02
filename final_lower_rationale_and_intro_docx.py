from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor


SRC = r"N:\개인\다음논문\LLM_DIF_full_manuscript_draft_psychometric_style_정리본_rationale낮춤.docx"
OUT = r"N:\개인\다음논문\LLM_DIF_full_manuscript_draft_psychometric_style_정리본_rationale낮춤_최종.docx"


def set_text(p, text):
    p.clear()
    r = p.add_run(text)
    r.font.name = "맑은 고딕"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
    r.font.size = Pt(10.5)
    r.bold = False
    r.font.color.rgb = RGBColor(0, 0, 0)


def restyle(p):
    txt = p.text.strip()
    for r in p.runs:
        r.font.name = "맑은 고딕"
        r._element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
        if txt.startswith(tuple(f"{i}." for i in range(1, 7))):
            r.font.size = Pt(11 if "." in txt[:4] else 13)
            r.bold = True
            r.font.color.rgb = RGBColor(31, 55, 99)
        elif txt.startswith("표 "):
            r.font.size = Pt(9.3)
            r.bold = True
            r.font.color.rgb = RGBColor(68, 68, 68)
        else:
            r.font.size = Pt(10.5)


def remove_paragraph(p):
    p._element.getparent().remove(p._element)


def main():
    doc = Document(SRC)

    for p in list(doc.paragraphs):
        txt = p.text.strip()

        if txt.startswith("문제는 MAPS와 같은 대규모 패널 자료에서 검토해야 할 문항-공변량 조합"):
            set_text(
                p,
                "문제는 MAPS와 같은 대규모 패널 자료에서 검토해야 할 문항-공변량 조합의 수가 빠르게 증가한다는 점이다. 문항 수와 주요 공변량 수가 늘어날수록 전문가가 검토해야 하는 단위는 개별 문항이 아니라 문항-공변량 조합으로 확장된다. 모든 조합에 대해 인간 전문가가 동일한 수준의 주의력을 유지하며 문항 내용을 읽고 잠재적인 DIF 작동 기제를 추론하기는 현실적으로 어렵다. 단순 키워드 매칭(keyword matching) 방식, 예컨대 문항 텍스트 내에 ‘차별’, ‘외국인’, ‘언어’와 같은 명시적 단어가 포함되어 있는지 여부로 선별하는 방식은 표면 단서가 분명한 일부 문항을 빠르게 거를 수 있으나, 어휘가 직접 겹치지 않으면서 인지적 복잡성이나 문화적 맥락을 경유해 간접적으로 작동하는 DIF 후보군을 포착하는 데는 한계가 있다.",
            )

        elif txt.startswith("이 지점에서 대규모 언어모형"):
            set_text(
                p,
                "이 지점에서 대규모 언어모형(large language model, 이하 LLM)은 DIF 판정 도구가 아니라, 통계적 분석 및 전문가 검토 이전 단계에서 작동하는 후보 가설 생성 및 우선순위화 도구로 검토될 수 있다. LLM은 문항 문장, 구인의 이론적 설명, 응답 범주, 공변량 정의를 통합적으로 입력받아 문맥적 정합성을 언어적으로 추론할 수 있다. 별도의 예시나 경험적 DIF 분석 결과를 제공하지 않은 조건에서 LLM이 각 문항-공변량 조합의 DIF 가능성, 예상 방향, 판단 근거를 가설 형태로 제시한다면, 연구자는 후속 통계 분석 결과와 대조할 검토 후보를 보다 정형화된 방식으로 정렬할 수 있다. 이때 LLM의 산출물은 최종적인 심리측정학적 증거가 아니라, 통계적·실증적 검증을 요구하는 후보 가설의 역할을 수행한다.",
            )

        elif txt.startswith("물론 LLM의 추론 결과를 심리측정 분야에 적용할 때"):
            set_text(
                p,
                "물론 LLM의 추론 결과를 심리측정 분야에 적용할 때는 엄격한 경계조건이 필요하다. LLM은 자연스러운 설명을 생성하지만, 잠재특성의 실제 집단 차이(impact)와 문항의 기능적 왜곡(DIF)을 개념적으로 혼동하거나, 성별·연령에 대한 일반론적인 사회적 고정관념을 그럴듯한 DIF 기제로 둔갑시킬 위험이 있다. 또한 폐쇄형 LLM의 경우 훈련 데이터에 심리측정학 문헌이 포함되어 있어 유사 맥락의 연구 흐름을 학습했을 가능성(data leakage)도 배제하기 어렵다. 따라서 본 연구는 LLM의 가능성을 무조건적으로 옹호하기보다, 실제 수집된 패널 자료에서 산출한 잠정적 순서형 DIF 선별 결과를 비교 기준으로 삼아 LLM의 판단이 단순 키워드 매칭 대비 어떤 추가 정보를 제공하는지 검토한다. 또한 기본 지시문과 엄격한 DIF 구분 지시문을 비교하여, 지시문 조건에 따라 LLM의 우선순위화 결과와 설명 양식이 어떻게 달라지는지 살펴본다.",
            )

        elif txt.startswith("본 연구는 MAPS 청소년 및 보호자 응답 문항을 대상으로, 실제 자료에서 도출된 경험적 DIF 탐색 결과"):
            set_text(
                p,
                "본 연구는 MAPS 청소년 및 보호자 응답 문항을 대상으로, 순서형 로지스틱 회귀에 기반한 잠정적 DIF 선별 결과를 비교 기준으로 삼아 LLM의 우선순위화 성능을 평가한다. 이 선별 결과는 최종 DIF 판정이나 참값이 아니라, LLM 산출물과 키워드 기준이 경험적 선별 신호와 어느 정도 정렬되는지 평가하기 위한 provisional criterion이다. 구체적인 연구 질문은 다음과 같다.",
            )

        elif txt.startswith("연구질문 2:"):
            set_text(
                p,
                "연구질문 2: LLM이 산출한 DIF 가능성 점수는 텍스트 내 지표 단어 포함 여부를 뜻하는 이분형 키워드 기준보다 잠정적 경험 DIF 후보를 더 효율적으로 우선순위화하는가?",
            )

        elif txt.startswith("연구질문 4:"):
            set_text(
                p,
                "연구질문 4: 실제 차이(impact)와 문항 기능 차이(DIF)의 구분을 명시적으로 제약한 엄격한 지시문(strict prompt)은 기본 지시문(base prompt)에 비해 LLM의 우선순위화 성능과 상위 후보 목록의 안정성을 개선하는가? 판단 근거 분석은 이 질문에 대한 타당성 검증이 아니라, 지시문 조건에 따라 LLM 설명에서 나타나는 보조적 failure-mode 신호가 어떻게 달라지는지 기술하기 위한 부가 분석으로 다룬다.",
            )

        elif txt == "다음 결과는 LLM 판단 근거의 타당성을 검증한 것이 아니라, 자동 규칙에 따른 설명 양식과 위험 신호의 분포를 기술한 것이다.":
            remove_paragraph(p)
            continue

        elif txt.startswith("주. RI = random-intercept ordinal probe. 각 사분면에서 대표 사례"):
            set_text(
                p,
                "주. RI = random-intercept ordinal probe. 각 사분면에서 focal 사례 1개를 선정하였으며, 응답자 수가 큰 경우 고정 seed로 최대 400명까지 표집하였다. 이 분석은 전체 검증이 아니라 focal triangulation이다.",
            )

        restyle(p)

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
