from pathlib import Path
import csv

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


ROOT = Path(".")
OUT_DIR = ROOT / "RESEARCH" / "llm_dif_paper_pipeline" / "outputs"
DATA_DIR = ROOT / "llm_dif_output"
MD_OUT = OUT_DIR / "method_results_discussion_revised_for_advisor_ko.md"
DOCX_OUT = OUT_DIR / "논문초고_방법결과논의_수정본.docx"


def read_csv(path):
    with path.open("r", encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


metrics = read_csv(DATA_DIR / "maps_llm_gemini_sensitivity_eval_metrics.csv")
metrics_w6 = read_csv(DATA_DIR / "maps_llm_gemini_sensitivity_eval_metrics_w6.csv")
rank_overlap = read_csv(DATA_DIR / "maps_llm_gemini_sensitivity_rank_overlap.csv")
rationale = read_csv(DATA_DIR / "maps_llm_sensitivity_rationale_code_summary.csv")


def f3(x):
    if x in ("", None):
        return ""
    try:
        v = float(x)
    except Exception:
        return str(x)
    s = f"{v:.3f}"
    return s[1:] if s.startswith("0") else s


def pct(n, d):
    return f3(float(n) / float(d))


def by_metric(rows, prompt, group):
    for row in rows:
        if row["prompt_version"] == prompt and row["group"] == group:
            return row
    raise KeyError((prompt, group))


cov_label = {
    "covariate:discrim_any": "차별 경험",
    "covariate:korean_c": "한국어 능력",
    "covariate:gender": "성별",
    "covariate:age_c": "연령",
    "covariate:income_c": "가구소득",
}

prompt_label = {
    "original": "기본 지시문",
    "strict_dif": "엄격한 DIF 구분 지시문",
}


def overall_rows(rows):
    out = []
    for prompt in ["original", "strict_dif"]:
        m = by_metric(rows, prompt, "overall")
        out.append([
            prompt_label[prompt],
            m["n_pairs"],
            m["positives"],
            pct(m["positives"], m["n_pairs"]),
            pct(m["positives"], m["n_pairs"]),
            f3(m["llm_average_precision"]),
            f3(m["keyword_average_precision"]),
            f3(m["llm_precision_at_5"]),
            f3(m["keyword_precision_at_5"]),
            f3(m["llm_precision_at_10"]),
            f3(m["keyword_precision_at_10"]),
        ])
    return out


def cov_rows(rows):
    out = []
    for group, label in cov_label.items():
        for prompt in ["original", "strict_dif"]:
            m = by_metric(rows, prompt, group)
            out.append([
                label,
                prompt_label[prompt],
                m["n_pairs"],
                m["positives"],
                pct(m["positives"], m["n_pairs"]),
                f3(m["llm_average_precision"]),
                f3(m["keyword_average_precision"]),
            ])
    return out


def rank_rows():
    out = []
    for r in rank_overlap:
        out.append([
            "기본 지시문 vs 엄격한 DIF 구분 지시문",
            r["common_pairs"],
            f"top-{r['k']} overlap",
            f3(r["topk_overlap"]),
        ])
    if rank_overlap:
        out.insert(0, [
            "기본 지시문 vs 엄격한 DIF 구분 지시문",
            rank_overlap[0]["common_pairs"],
            "Spearman 순위상관",
            f3(rank_overlap[0]["spearman_score_cor"]),
        ])
    return out


def rationale_rows():
    labels = {
        "RESPONSE_PROCESS": "응답 과정 근거",
        "TESTABLE_HYPOTHESIS": "검증 가능한 가설 표현",
        "NO_WITHIN_TRAIT_CONDITIONING": "같은 잠재특성 조건 누락",
        "VAGUE_GENERALITY": "모호한 일반론",
        "STEREOTYPE_GENDER_AGE": "성별·연령 고정관념 가능성",
        "ITEM_IRRELEVANT_SPECULATION": "문항 관련성이 약한 추측",
        "IMPACT_DIF_CONFUSION": "실제 차이를 DIF로 오인할 가능성",
    }
    wanted = list(labels.keys())
    by_cat = {}
    for r in rationale:
        if r.get("level") == "prompt":
            by_cat[(r["prompt_version"], r["code"])] = r
    out = []
    for cat in wanted:
        orig = by_cat.get(("original", cat), {})
        strict = by_cat.get(("strict_dif", cat), {})
        out.append([
            labels[cat],
            f3(orig.get("rate", "")),
            f3(strict.get("rate", "")),
        ])
    return out


overall = overall_rows(metrics)
covariates = cov_rows(metrics)
w6_overall = overall_rows(metrics_w6)
w6_covariates = cov_rows(metrics_w6)
rank_table = rank_rows()
rationale_table = rationale_rows()


def md_table(headers, rows):
    lines = ["| " + " | ".join(headers) + " |"]
    aligns = []
    for h in headers:
        aligns.append("---:" if any(k in h for k in ["수", "AP", "P@", "비율", "값"]) else "---")
    lines.append("| " + " | ".join(aligns) + " |")
    for row in rows:
        lines.append("| " + " | ".join(str(x) for x in row) + " |")
    return "\n".join(lines)


title = "대규모 언어모형을 활용한 차별기능문항 후보 가설 생성의 가능성과 한계: 다문화청소년패널 문항-공변량 조합의 우선순위화 분석"

abstract = (
    "본 연구는 다문화청소년패널조사(MAPS) 2기 자료의 심리사회적 문항을 대상으로, "
    "생성형 대규모 언어모형이 차별기능문항(DIF) 검토를 위한 문항-공변량 후보 가설을 "
    "우선순위화할 수 있는지 탐색하였다. 분석 단위는 개별 응답자가 아니라 문항-공변량 "
    "조합이었다. Gemini 2.5 Flash에는 문항 문장, 구인 맥락, 응답 범주, 공변량 정의를 "
    "제공하되 잠정적 경험 DIF 선별 결과는 제공하지 않았다. 기본 지시문과 엄격한 DIF "
    "구분 지시문을 비교하였고, LLM 점수의 우선순위화 성능은 문항 텍스트 기반 키워드 "
    "비교 기준 및 wave를 통제한 잠정적 ordinal DIF 선별 결과와 비교하였다. 1-5차년도 "
    "pooled long 자료에서 기본 지시문 LLM의 AP는 .382, 엄격한 DIF 구분 지시문 LLM의 "
    "AP는 .405였으며, 두 조건 모두 키워드 비교 기준의 AP(.441, .445)를 넘지 못했다. "
    "6차년도 단일 wave 민감도 분석에서도 LLM은 키워드 기준을 일관되게 능가하지 못했다. "
    "엄격한 지시문은 일부 판단 근거 지표를 개선했지만, 상위 후보 목록은 지시문 조건에 "
    "민감했고 성별·연령 관련 일반화 가능성도 남아 있었다. 따라서 LLM 산출물은 DIF "
    "판정 근거가 아니라 키워드 기준, 경험적 선별, 전문가 검토, 후속 심리측정 모형과 "
    "결합해 제한적으로 사용할 수 있는 후보 가설 자료로 해석되어야 한다."
)

method_text = [
    (
        "본 연구는 MAPS 2기 1-5차년도 자료를 주 분석 자료로 사용하였다. 자료는 청소년과 "
        "보호자를 반복 추적한 종단 패널이지만, 본 연구의 목적은 종단 측정동일성 또는 "
        "MNLFA를 직접 추정하는 데 있지 않다. 분석의 초점은 경험적 선별 이후 연구자가 "
        "검토할 문항-공변량 후보를 어떻게 우선순위화할 수 있는지에 있다. 따라서 1-5차년도 "
        "자료는 문항 응답과 공변량 정보를 결합하기 위해 pooled long format으로 구성하였고, "
        "6차년도 자료는 단일 wave 민감도 분석에 사용하였다."
    ),
    (
        "최종 문항 pool은 105개 문항으로 구성하였다. 보호자 문항에는 보호자 공변량을, "
        "청소년 문항에는 청소년 공변량을 적용하였다. 공변량은 차별 경험, 한국어 능력, "
        "가구소득, 성별, 연령이었다. 다만 성별 공변량은 청소년 문항에만 적용하였고, "
        "응답자 유형상 부적절한 문항-공변량 조합은 제외하였다. 이 절차를 거쳐 기본 "
        "지시문 조건에서는 486개, 엄격한 DIF 구분 지시문 조건에서는 482개 문항-공변량 "
        "조합이 평가에 포함되었다."
    ),
    (
        "LLM에는 Gemini 2.5 Flash를 사용하였다. API 호출에서는 temperature를 0으로 설정하고 "
        "응답 형식은 JSON으로 제한하였다. 입력에는 문항 문장, 구인 정보, 응답 범주, "
        "공변량 정의를 포함하였다. 경험적 DIF 선별 결과, p값, 효과크기, 문항 통계량은 "
        "입력하지 않았다. 따라서 LLM 산출물은 경험적 결과를 보지 않은 상태에서 생성된 "
        "후보 가설로 간주하였다. 모든 raw response와 파싱 결과는 별도로 보존하였다."
    ),
    (
        "비교 기준으로 키워드 비교 기준을 사용하였다. 이 기준은 문항 텍스트에 공변량 관련 "
        "핵심어가 포함되는지를 사전에 정한 규칙에 따라 점수화한 것이다. 키워드 비교 기준은 "
        "LLM 출력, 응답 분포, 문항 통계량, 경험적 DIF label을 사용하지 않는다. 다만 이 기준은 "
        "단순한 무작위 기준이 아니라 연구자가 구성한 lexical benchmark이므로, 차별 경험이나 "
        "한국어 능력처럼 표면 단서가 뚜렷한 공변량에서는 강한 비교 기준으로 작동할 수 있다."
    ),
    (
        "잠정적 경험 DIF 선별은 누적로짓 proportional odds 모형을 이용하였다. 각 문항 응답은 "
        "순서형 결과변수로 두었고, 예측변수에는 leave-one-item-out 방식으로 구성한 척도 점수 "
        "proxy, 해당 공변량, wave를 포함하였다. 즉 각 문항에 대해 '문항 응답 = 척도 점수 proxy + "
        "공변량 + wave' 형태의 screening 모형을 적합하였다. 공변량 계수는 같은 척도 점수 수준에서 "
        "공변량에 따라 문항 응답 경향이 달라지는지를 보는 uniform ordinal DIF 예비 신호로 해석하였다."
    ),
    (
        "이 선별 결과는 최종 DIF 판정이나 참값이 아니다. 본 연구에서는 p값에 BH 방식의 FDR 보정을 "
        "적용하고, 절대 계수값이 .20 이상인 경우를 practical nonzero 조건으로 두었다. FDR 보정 "
        "p값이 .05 미만이고 practical nonzero 조건을 만족한 조합을 잠정적 경험 DIF 후보로 표시하였다. "
        "이 기준은 LLM과 키워드 비교 기준의 우선순위화 성능을 평가하기 위한 provisional label로만 사용하였다."
    ),
    (
        "주요 평가지표는 average precision(AP)이었다. AP는 경험적 후보가 점수 순위의 상위에 얼마나 "
        "잘 배치되는지를 나타낸다. AP는 양성 후보 비율의 영향을 받으므로, 결과표에는 positive prevalence를 "
        "무작위 순위화의 기준 AP로 함께 제시하였다. 보조 지표로 precision@5와 precision@10을 보고하였다. "
        "또한 기본 지시문과 엄격한 DIF 구분 지시문 간의 민감도를 확인하기 위해 Spearman 순위상관과 상위 후보 overlap을 산출하였다."
    ),
]

discussion_text = [
    (
        "첫째, LLM은 문항-공변량 조합별 후보 가설을 구조화된 형식으로 생성할 수 있었다. "
        "두 지시문 조건 모두 대부분의 응답이 JSON 형식으로 파싱되었고, 문항과 공변량 사이의 "
        "가능한 연결을 판단 근거로 제시하였다. 그러나 이 결과는 LLM이 DIF를 판정할 수 있음을 "
        "의미하지 않는다. LLM 산출물은 경험적 검증을 기다리는 후보 가설에 머물러야 한다."
    ),
    (
        "둘째, LLM은 키워드 비교 기준을 일관되게 능가하지 못했다. 주 분석에서 기본 지시문과 "
        "엄격한 DIF 구분 지시문의 LLM AP는 각각 .382와 .405였고, 키워드 비교 기준의 AP는 "
        "각각 .441과 .445였다. 특히 한국어 능력과 차별 경험에서는 표면 단서가 강하게 작동하여 "
        "키워드 기준이 더 높은 성능을 보였다. 따라서 본 연구의 결과는 LLM의 일반적 우수성을 "
        "보여주기보다, 단순하고 명시적인 비교 기준을 반드시 함께 두어야 함을 보여준다."
    ),
    (
        "셋째, 성능은 공변량에 따라 달랐다. 차별 경험에서는 엄격한 지시문이 LLM AP를 키워드 "
        "기준에 가깝게 높였지만 넘지는 못했다. 성별에서는 LLM이 키워드 기준보다 높은 AP를 "
        "보였으나, 후보 수와 판단 근거의 성격을 고려하면 강한 결론으로 해석하기 어렵다. "
        "가구소득은 경험 후보 수가 매우 적어 AP가 불안정하였다. 공변량별 결과는 LLM의 성능을 "
        "전체 평균만으로 판단하기 어렵다는 점을 보여준다."
    ),
    (
        "넷째, 엄격한 DIF 구분 지시문은 일부 지표를 개선했지만 충분한 해결책은 아니었다. "
        "전체 AP와 일부 판단 근거 코딩 지표는 개선되었지만, P@5와 P@10은 오히려 낮아졌고 "
        "상위 후보 목록은 기본 지시문과 거의 겹치지 않았다. 이는 LLM 기반 후보 생성이 지시문 "
        "조건에 민감하며, 단일 지시문 결과를 안정적인 후보 목록으로 간주하기 어렵다는 점을 시사한다."
    ),
    (
        "본 연구의 한계도 분명하다. 잠정적 ordinal DIF 선별 결과는 참값이 아니며, screening "
        "모형의 가정과 기준에 따라 달라질 수 있다. 1-5차년도 pooled long 분석은 wave를 통제했지만, "
        "동일 응답자의 반복 관측에 따른 의존성을 완전히 해결한 것은 아니다. 따라서 경험적 label은 "
        "anti-conservative할 수 있으며, 후속 연구에서는 cluster-robust 표준오차, GEE, random intercept "
        "모형 또는 응답자 단위 부트스트랩을 고려할 필요가 있다."
    ),
    (
        "또한 본 연구는 단일 LLM인 Gemini 2.5 Flash와 두 지시문 조건에 기반한다. 폐쇄형 LLM은 "
        "모델 버전과 API 설정이 시간에 따라 바뀔 수 있으므로 재현성에 제한이 있다. 판단 근거 코딩도 "
        "자동 규칙 기반이므로 LLM 설명의 실제 타당성을 입증하지 않는다. 일부 표본에 대해서는 전문가 "
        "코딩을 통해 판단 근거의 심리측정학적 적절성을 검증할 필요가 있다."
    ),
    (
        "결론적으로, 본 연구는 LLM을 DIF 판정 도구로 지지하지 않는다. 오히려 LLM은 키워드 기준과 "
        "경험적 선별 결과에 비추어 검토해야 할 후보 가설 생성 도구로 제한적으로 사용될 때 의미가 있다. "
        "LLM의 가치는 답을 내리는 데 있지 않고, 어떤 문항-공변량 조합을 먼저 검토할지 제안하고 그 제안이 "
        "어디서 심리측정 검증과 맞거나 어긋나는지를 드러내는 데 있다."
    ),
]


sections = []
sections.append(f"# {title}\n")
sections.append("## 초록\n\n" + abstract + "\n")
sections.append("## 연구방법\n\n" + "\n\n".join(method_text) + "\n")
sections.append("### 표 1. 주 분석 전체 우선순위화 성능\n\n" + md_table(
    ["지시문 조건", "분석 조합 수", "경험 후보 수", "positive prevalence", "baseline AP", "LLM AP", "키워드 AP", "LLM P@5", "키워드 P@5", "LLM P@10", "키워드 P@10"],
    overall,
) + "\n")
sections.append("### 표 2. 공변량별 우선순위화 성능\n\n" + md_table(
    ["공변량", "지시문 조건", "분석 조합 수", "경험 후보 수", "positive prevalence", "LLM AP", "키워드 AP"],
    covariates,
) + "\n")
sections.append("### 표 3. 지시문 민감도\n\n" + md_table(
    ["비교", "공통 조합 수", "지표", "값"],
    rank_table,
) + "\n")
sections.append("### 표 4. 판단 근거 자동 코딩 결과\n\n" + md_table(
    ["코딩 범주", "기본 지시문", "엄격한 DIF 구분 지시문"],
    rationale_table,
) + "\n")
sections.append("### 표 5. 6차년도 단일 wave 민감도 분석: 전체 성능\n\n" + md_table(
    ["지시문 조건", "분석 조합 수", "경험 후보 수", "positive prevalence", "baseline AP", "LLM AP", "키워드 AP", "LLM P@5", "키워드 P@5", "LLM P@10", "키워드 P@10"],
    w6_overall,
) + "\n")
sections.append("### 표 6. 6차년도 단일 wave 민감도 분석: 공변량별 성능\n\n" + md_table(
    ["공변량", "지시문 조건", "분석 조합 수", "경험 후보 수", "positive prevalence", "LLM AP", "키워드 AP"],
    w6_covariates,
) + "\n")
sections.append("## 논의 및 결론\n\n" + "\n\n".join(discussion_text) + "\n")

MD_OUT.write_text("\n".join(sections), encoding="utf-8")


def set_font(run, size=10.5, bold=False):
    run.font.name = "맑은 고딕"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
    run.font.size = Pt(size)
    run.bold = bold


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def add_para(doc, text, indent=True):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.45
    p.paragraph_format.space_after = Pt(4)
    if indent:
        p.paragraph_format.first_line_indent = Cm(0.7)
    run = p.add_run(text)
    set_font(run)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        set_font(run, size=14 if level == 1 else 12, bold=True)
    return p


def set_cell_text(cell, text, bold=False):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(str(text))
    set_font(run, size=8.5, bold=bold)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(doc, title_text, headers, rows):
    caption = doc.add_paragraph()
    caption.paragraph_format.space_before = Pt(6)
    caption.paragraph_format.space_after = Pt(3)
    run = caption.add_run(title_text)
    set_font(run, size=10, bold=True)

    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], h, bold=True)
        set_cell_shading(table.rows[0].cells[i], "EDEDED")
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_text(cells[i], value)
    doc.add_paragraph()


doc = Document()
section = doc.sections[0]
section.top_margin = Cm(2.2)
section.bottom_margin = Cm(2.0)
section.left_margin = Cm(2.2)
section.right_margin = Cm(2.2)

style = doc.styles["Normal"]
style.font.name = "맑은 고딕"
style._element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
style.font.size = Pt(10.5)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run(title)
set_font(run, size=15, bold=True)
doc.add_paragraph()

add_heading(doc, "초록", level=1)
add_para(doc, abstract)

add_heading(doc, "연구방법", level=1)
for t in method_text:
    add_para(doc, t)

add_heading(doc, "결과", level=1)
add_table(doc, "표 1. 주 분석 전체 우선순위화 성능",
          ["지시문 조건", "분석 조합 수", "경험 후보 수", "positive prevalence", "baseline AP", "LLM AP", "키워드 AP", "LLM P@5", "키워드 P@5", "LLM P@10", "키워드 P@10"],
          overall)
add_table(doc, "표 2. 공변량별 우선순위화 성능",
          ["공변량", "지시문 조건", "분석 조합 수", "경험 후보 수", "positive prevalence", "LLM AP", "키워드 AP"],
          covariates)
add_table(doc, "표 3. 지시문 민감도",
          ["비교", "공통 조합 수", "지표", "값"],
          rank_table)
add_table(doc, "표 4. 판단 근거 자동 코딩 결과",
          ["코딩 범주", "기본 지시문", "엄격한 DIF 구분 지시문"],
          rationale_table)
add_table(doc, "표 5. 6차년도 단일 wave 민감도 분석: 전체 성능",
          ["지시문 조건", "분석 조합 수", "경험 후보 수", "positive prevalence", "baseline AP", "LLM AP", "키워드 AP", "LLM P@5", "키워드 P@5", "LLM P@10", "키워드 P@10"],
          w6_overall)
add_table(doc, "표 6. 6차년도 단일 wave 민감도 분석: 공변량별 성능",
          ["공변량", "지시문 조건", "분석 조합 수", "경험 후보 수", "positive prevalence", "LLM AP", "키워드 AP"],
          w6_covariates)

add_heading(doc, "논의 및 결론", level=1)
for t in discussion_text:
    add_para(doc, t)

doc.save(DOCX_OUT)
print(MD_OUT.resolve())
print(DOCX_OUT.resolve())
