from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph


SRC = Path(r"C:\chen_bauer_2024\docx_work\submission_trim_input.docx")
TMP = Path(r"C:\chen_bauer_2024\docx_work\submission_trim_step1.docx")
OUT = Path(r"C:\chen_bauer_2024\docx_work\submission_trim.docx")


def delete_element(el) -> None:
    parent = el.getparent()
    if parent is not None:
        parent.remove(el)


def set_text(p, text: str) -> None:
    for run in p.runs:
        run.text = ""
    if p.runs:
        p.runs[0].text = text
    else:
        p.add_run(text)


def insert_before(paragraph, text: str, style=None) -> Paragraph:
    new_p = OxmlElement("w:p")
    paragraph._p.addprevious(new_p)
    p = Paragraph(new_p, paragraph._parent)
    if style is not None:
        p.style = style
    p.add_run(text)
    return p


def insert_table_before(doc: Document, paragraph: Paragraph, rows: list[list[str]]) -> None:
    tbl = OxmlElement("w:tbl")
    # Build table through a temporary document, then move XML.
    tmp = Document()
    t = tmp.add_table(rows=len(rows), cols=len(rows[0]))
    t.style = "Table Grid"
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            t.cell(r, c).text = val
    paragraph._p.addprevious(t._tbl)


# Step 1: remove calibration paragraph/table and retitle section 3.4.
doc = Document(SRC)
for p in doc.paragraphs:
    if p.text.strip() == "3.4 응답자 유형별 성능과 LLM 점수 보정":
        set_text(p, "3.4 응답자 유형별 성능")
    elif p.text.strip() == "표 9. 지시문 조건 간 민감도":
        set_text(p, "표 8. 지시문 조건 간 민감도")

for p in list(doc.paragraphs):
    if p.text.strip().startswith("LLM 점수의 보정 가능성도 별도로 확인하였다"):
        delete_element(p._element)
    elif p.text.strip() == "표 8. LLM 점수 구간별 선별 양성률":
        delete_element(p._element)

# Remove the calibration table: header begins with "지시문 조건 | 점수 구간".
for t in list(doc.tables):
    if t.rows and len(t.rows[0].cells) >= 5:
        header = " | ".join(cell.text for cell in t.rows[0].cells)
        if "점수 구간" in header and "평균 LLM 점수" in header:
            delete_element(t._element)
            break

doc.save(TMP)

# Step 2: remove detailed sections 3.6-3.12 and insert one concise section.
doc = Document(TMP)
body = doc.element.body
start_el = None
end_el = None
for p in doc.paragraphs:
    txt = p.text.strip()
    if txt.startswith("3.6 보조 오류 유형 점검 결과"):
        start_el = p._element
    if txt.strip() == "4. 논의":
        end_el = p._element
        break

if start_el is not None and end_el is not None:
    deleting = False
    for el in list(body):
        if el is start_el:
            deleting = True
        if el is end_el:
            deleting = False
        if deleting:
            delete_element(el)

# Reopen after XML deletion.
doc.save(TMP)
doc = Document(TMP)
discussion = None
for p in doc.paragraphs:
    if p.text.strip() == "4. 논의":
        discussion = p
        break
assert discussion is not None

insert_before(discussion, "3.6 민감도 및 보조 분석 요약", style=doc.styles["Heading 3"])
insert_before(
    discussion,
    "주 분석에서 나타난 Gemini의 소폭 AP 우위는 추가 분석에서 안정적인 차이로 확인되지 않았다. 문항 단위 cluster bootstrap 구간은 두 지시문 조건 모두에서 0을 포함했고, paired permutation test에서도 유의한 차이가 나타나지 않았다. 6차년도 단일 wave 분석에서는 오히려 전문가 어휘 기준의 AP가 Gemini보다 높았다. TF-IDF n-gram 의미 유사도 기준은 Gemini보다 낮은 AP를 보였으나, 이는 Gemini가 전문가 어휘 기준을 안정적으로 능가했다는 결론과는 별개의 결과이다.",
    style=doc.styles["Normal"],
)
insert_before(
    discussion,
    "보조 분석은 LLM 산출물의 사용 조건을 더 구체화한다. 기본 지시문과 엄격 지시문의 top-5 및 top-10 후보는 서로 겹치지 않아 상위 후보 목록이 지시문에 민감함을 보였다. 두 기준이 동시에 높은 both-high 조합은 경험적 양성률이 가장 높았지만, LLM 단독 고득점 조합은 그만큼 강한 후보군을 만들지 못했다. 판단 근거 자동 점검은 설명의 타당성을 입증하기보다, 성별·연령 일반화나 impact-DIF 혼동 같은 위험 신호를 확인하는 보조 절차로 해석하였다.",
    style=doc.styles["Normal"],
)
insert_before(discussion, "표 9. 주요 민감도 및 보조 분석 요약", style=doc.styles["Caption"])
insert_table_before(
    doc,
    discussion,
    [
        ["분석", "핵심 결과", "해석"],
        ["불확실성 평가", "AP 차이의 bootstrap 95% 구간이 0을 포함; permutation test도 비유의", "Gemini의 AP 우위는 안정적 차이로 보기 어렵다."],
        ["6차년도 단일 wave", "기본/엄격 지시문 모두 전문가 어휘 기준 AP가 Gemini보다 높음", "pooled 분석의 소폭 우위가 단일 시점에서 재현되지 않았다."],
        ["TF-IDF 기준", "TF-IDF n-gram AP는 Gemini보다 낮음", "Gemini는 단순 문자열 유사도보다 강하지만, 전문가 어휘 기준을 안정적으로 넘지는 못했다."],
        ["사분면 분석", "both-high 양성률이 가장 높음(.538/.636)", "LLM과 전문가 어휘 기준이 수렴하는 조합을 우선 검토하는 전략이 가장 방어 가능하다."],
        ["판단 근거 점검", "엄격 지시문은 일부 설명 양식을 바꾸었으나 오류 신호를 제거하지 못함", "설명은 타당성 증거가 아니라 검토해야 할 가설 자료이다."],
    ],
)
insert_before(
    discussion,
    "주. 세부 민감도 표, 판단 근거 사례, random-intercept probe 결과는 투고 시 부록 또는 온라인 보충자료로 제시하는 것이 적절하다.",
    style=doc.styles["Normal"],
)

# Rename discussion heading style if it was accidentally caption in previous drafts.
for p in doc.paragraphs:
    if p.text.strip() == "4. 논의":
        p.style = doc.styles["Heading 2"]
    elif p.text.strip() in {"4.1 주요 발견 및 논의", "4.2 학술적 기여 및 실무적 제언", "4.3 해석 범위와 후속 과제"}:
        p.style = doc.styles["Heading 3"]

# Clean obvious stale references to removed table numbers in result/discussion prose.
for p in doc.paragraphs:
    text = p.text
    new = text.replace("표 19", "표 9")
    new = new.replace("표 18", "표 9")
    new = new.replace("3.12절 사분면 분석 참조", "3.6절 참조")
    if new != text:
        set_text(p, new)

doc.save(OUT)
print(OUT)
