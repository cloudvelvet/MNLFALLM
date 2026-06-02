from __future__ import annotations

from pathlib import Path

from docx import Document


SRC = Path(r"C:\chen_bauer_2024\docx_work\senior_style_revision_clean2.docx")
OUT = Path(r"C:\chen_bauer_2024\docx_work\senior_style_revision_clean3.docx")


def set_text(p, text: str) -> None:
    for run in p.runs:
        run.text = ""
    if p.runs:
        p.runs[0].text = text
    else:
        p.add_run(text)


doc = Document(SRC)

set_text(
    doc.paragraphs[69],
    "각 영역에 대해서는 주 선별 기준과 cluster-robust 선별 기준의 양성 비율을 비교하였다. 추가로, 각 지시문 조건에서 네 사분면별로 1개 문항-공변량 조합을 선정하여 총 8개 조합에 대해 random-intercept ordinal probe를 수행하였다. 이 사례들은 각 사분면 전체를 대표한다고 가정하기 위한 것이 아니라, 사분면별 후보가 응답자별 무선절편을 포함한 모형에서도 같은 방향의 신호를 보이는지 점검하기 위한 표적 사례로 선정하였다. 선정은 경험적 선별 결과를 새로 최적화하는 방식이 아니라, 해당 사분면에 속하면서 분석 사례 수가 충분하고 모형 수렴이 가능한 조합을 우선하는 규칙에 따라 이루어졌다. 이 probe는 전체 MNLFA가 아니라 표적 경험 검토이며, 계수 방향과 선별 label이 모형 구조 변화에 따라 얼마나 안정적인지 확인하기 위한 보조 분석이다.",
)

set_text(
    doc.paragraphs[119],
    "그러나 효과크기 임계치를 |.30| 이상으로 엄격하게 설정하여 DIF 양성 비율이 .177로 떨어지는 조건에서는, 기본 지시문 조건의 LLM AP(.326)가 전문가 어휘 기준(.299)을 다소 상회하는 경향을 보였다. 반면, 효과크기 임계치를 적용하지 않고 FDR p값만으로 선별한 조건(양성 비율 .490)에서는 전문가 어휘 기준의 AP(.573)가 LLM AP(.553)보다 높게 나타났다.",
)

set_text(
    doc.paragraphs[156],
    "셋째, 성능은 공변량의 특성과 응답자 유형에 따라 달랐다. 차별 경험과 같이 문항 내용이 공변량의 의미와 직접 맞닿아 있는 경우 엄격 지시문 조건에서 Gemini AP가 전문가 어휘 기준을 상회하였다. 반면 한국어 능력이나 연령처럼 문항 내 표면 어휘와 공변량의 일치가 빈번한 경우에는 전문가 어휘 기준이 더 강하게 작동하였다. 응답자 유형별로는 청소년 문항보다 보호자 문항에서 Gemini의 AP 이점이 크게 나타났지만, 이 결과 역시 선별 기준과 문항 pool 구성에 의존한다. 특히 가구소득은 양성 후보 수가 매우 적어 공변량별 AP를 안정적인 성능 추정치로 해석하기 어렵다.",
)

for row in doc.tables[0].rows:
    for cell in row.cells:
        cell.text = cell.text.replace("wave를 통제한 잠정적 순서형 DIF 선별", "wave를 통제한 순서형 DIF 선별")

for p in doc.paragraphs:
    if "자동 audit 결과" in p.text:
        set_text(p, p.text.replace("자동 audit 결과", "자동 점검 결과"))

doc.save(OUT)
print(OUT)
