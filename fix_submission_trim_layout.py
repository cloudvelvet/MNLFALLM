from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK
from docx.text.paragraph import Paragraph
from docx.oxml import OxmlElement


WORK = Path(r"C:\chen_bauer_2024\docx_work")
SRC = WORK / "submission_trim.docx"
OUT = WORK / "submission_trim_final.docx"


def insert_paragraph_before(paragraph: Paragraph) -> Paragraph:
    new_p = OxmlElement("w:p")
    paragraph._p.addprevious(new_p)
    return Paragraph(new_p, paragraph._parent)


doc = Document(SRC)

target = "표 9. 주요 민감도 및 보조 분석 요약"
inserted = False
for paragraph in doc.paragraphs:
    if paragraph.text.strip() == target:
        page_break_paragraph = insert_paragraph_before(paragraph)
        page_break_paragraph.add_run().add_break(WD_BREAK.PAGE)
        inserted = True
        break

if not inserted:
    raise RuntimeError(f"Could not find caption paragraph: {target}")

doc.save(OUT)
print(OUT)
