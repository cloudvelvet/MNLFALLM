from pathlib import Path
import shutil

from docx import Document


WORK = Path(r"C:\chen_bauer_2024\docx_work")
SRC = WORK / "submission_trim_humanized_terms_repaired.docx"
LOCAL_OUT = WORK / "submission_trim_humanized_terms_final.docx"
FINAL_OUT = Path(r"N:\개인\다음논문\LLM_DIF_full_manuscript_draft_psychometric_style_Gemini범위수정_투고압축본_최종.docx")


REPLACEMENTS = [
    ("provisional criterion", "예비 평가 기준"),
    ("참값(gold standard)", "참값 기준"),
    ("단일 wave", "단일 조사시점"),
    ("단일 Wave", "단일 조사시점"),
    ("pooled 자료", "통합 자료"),
    ("pooled 분석", "통합 분석"),
    ("1-5차년도 pooled", "1-5차년도 통합"),
    ("criterion 정의", "선별 기준 정의"),
    ("top-k", "상위 후보군"),
    ("top-5", "상위 5개"),
    ("top-10", "상위 10개"),
    ("top-20", "상위 20개"),
]


def replace_in_paragraph(paragraph):
    if not paragraph.runs:
        return
    full = "".join(run.text for run in paragraph.runs)
    new = full
    for old, repl in REPLACEMENTS:
        new = new.replace(old, repl)
    if new == full:
        return
    first = paragraph.runs[0]
    for run in paragraph.runs[1:]:
        run.text = ""
    first.text = new


doc = Document(SRC)
for paragraph in doc.paragraphs:
    replace_in_paragraph(paragraph)
for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                replace_in_paragraph(paragraph)

doc.save(LOCAL_OUT)
shutil.copy2(LOCAL_OUT, FINAL_OUT)
print(LOCAL_OUT)
print(FINAL_OUT)
